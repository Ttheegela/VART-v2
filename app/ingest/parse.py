"""Files to numbered lines (spec 6.4). One paragraph, heading, list item or table row per line; table and
spreadsheet rows become record lines through app.text.record_line. The line rules mirror datakit/extract.py,
the reference the dev pack was validated with (tests/test_ingest_parse.py compares the two). PDFs live in
app/ingest/pdf.py. Limits are spec section 9's."""

import contextlib
import csv
import io
import re
import zipfile
from collections.abc import Iterator
from datetime import date, datetime
from itertools import zip_longest
from typing import Any

import docx
import openpyxl
from docx.enum.style import WD_STYLE_TYPE
from docx.styles import BabelFish
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.contracts import Line, LineKind, ParsedDocument
from app.text import cell_text, normalize, record_line

MAX_BYTES = 4 * 1024 * 1024  # spec 9: 4 MB per file (Vercel's request limit is 4.5 MB)
MAX_LINES = 20_000  # spec 9: per workspace, so also per document
MAX_LINE_CHARS = 20_000  # a line is never split, so one far longer could overflow its chunk's 1 MB tsvector
MAX_UNZIPPED = 50 * 1024 * 1024  # an Office file larger than this once unzipped is refused (zip bomb)
MAX_MEMBERS = 5_000
MAX_ROWS = 1_048_576  # rows read from one file, empty ones included (an Excel sheet's own limit)
MAX_CELLS = 5_000_000  # cells read from one file: a row naming column XFD is padded to 16,384 cells
MAX_DOCX_PART = (
    10 * 1024 * 1024
)  # word/document.xml and word/styles.xml, unzipped: python-docx holds them as trees
MAX_TABLE_COLUMNS = 63  # Word's own limit
FORMULA_NOTE = "(formula without a saved value)"
_TABLE_RULE = re.compile(r"^\|?\s*:?-{3,}")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_AS_OF = re.compile(r"\b(?:as of|as at|last updated)\b\W*(\d{4}-\d{2}-\d{2})", re.IGNORECASE)
_TO_COME = re.compile(
    r"\b(?:due|next|expir\w*|until|planned|target)\b", re.IGNORECASE
)  # "Next review due", "Expires"
_TEXT_FORMATS = ("csv", "md", "txt")
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # legacy .doc/.xls, and every password-protected Office file


class IngestError(ValueError):
    """A file the app will not read; the message is shown to the visitor as is."""


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def text_lines(text: str) -> list[Line]:
    """Markdown or plain text: '#' headings, '- '/'* ' bullets and pipe tables are their own lines; the other
    lines of a paragraph are joined (wrapping is not a paragraph break). A table row is dated by its latest
    date (spec 6.4)."""
    out: list[Line] = []
    para: list[str] = []
    header: list[str] | None = None

    def flush() -> None:
        if para:
            out.append(Line(normalize(" ".join(para))))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            header = None
            continue
        if line.lstrip().startswith("|"):
            flush()
            if header is None:
                header = _cells(line)
            elif not _TABLE_RULE.match(line.strip()):
                cells = _cells(line)
                out.append(Line(record_line(header, cells), "record", _latest_date(header, cells)))
            continue
        if line.startswith("#"):
            flush()
            out.append(Line(normalize(line.lstrip("#")), "heading"))
        elif line.lstrip().startswith(("- ", "* ")):
            flush()
            item = line.lstrip()[2:]
            if item.startswith(("[x] ", "[X] ")):  # a done task is a plain statement (Plan 2A Ruling 10)
                item = item[4:]
            out.append(Line(normalize(item)))
        else:
            para.append(line)
    flush()
    return [x for x in out if x.text]


def _table_rows(table: Table) -> Iterator[list[str]]:
    """The rows' cell texts, one pass over each row's w:tc elements. A cell spanning columns is repeated,
    never beyond the table's grid; a cell merged with the one above takes its text (python-docx's own
    `cells` walks both without a bound)."""
    tbl = table._tbl
    limit = min(len(tbl.tblGrid.gridCol_lst) or MAX_TABLE_COLUMNS, MAX_TABLE_COLUMNS)
    above: list[str] = []
    for tr in tbl.tr_lst:
        texts: list[str] = []
        for tc in tr.tc_lst:
            if tc.vMerge == "continue":
                text = above[len(texts)] if len(texts) < len(above) else ""
            else:
                text = _Cell(tc, table).text
            texts += [text] * min(tc.grid_span, max(limit - len(texts), 1))
        above = texts
        yield texts


def _docx_lines(data: bytes) -> list[Line]:
    document = docx.Document(io.BytesIO(data))
    if document.element.xpath("count(.//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo)"):
        # python-docx reads neither the inserted nor the deleted runs: the text left says what nobody wrote
        raise IngestError("This document has tracked changes. Accept or reject them, then upload it again.")
    # Styles resolve once: python-docx's Paragraph.style searches every style for every paragraph.
    names: dict[str, str] = {}
    for s in document.styles.element.style_lst:
        if s.type == WD_STYLE_TYPE.PARAGRAPH:
            names.setdefault(s.styleId, BabelFish.internal2ui(s.name_val) if s.name_val else "")
    default = document.styles.default(WD_STYLE_TYPE.PARAGRAPH)
    default_name = (default.name if default is not None else "") or ""
    out: list[Line] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                style = names.get(block._p.style, default_name) if block._p.style else default_name
                kind: LineKind = "heading" if style.startswith(("Heading", "Title")) else "text"
                out.append(Line(normalize(block.text), kind))
        elif isinstance(block, Table):
            rows = _table_rows(block)  # lazily: a hostile table has millions
            header = next(rows, [])
            for r in rows:
                out.append(Line(record_line(header, r), "record", _latest_date(header, r)))
                if len(out) > MAX_LINES:
                    break
        if len(out) > MAX_LINES:
            raise IngestError(f"This file has more than {MAX_LINES:,} lines.")
    return [x for x in out if x.text]


def _latest_date(header: list[Any], values: list[Any]) -> date | None:
    """The row's latest date. A cell under a header that names a date still to come ("Next review due",
    "Expires") is left out: it would date the row in the future (adversary checkpoint 1, M12)."""
    values = [v for h, v in zip_longest(header, values) if not _TO_COME.search(cell_text(h))]
    found = [v.date() if isinstance(v, datetime) else v for v in values if isinstance(v, date)]
    for v in values:
        if isinstance(v, str) and _ISO_DATE.match(v.strip()):
            with contextlib.suppress(ValueError):  # 2026-13-45 looks like a date and is not one
                found.append(date.fromisoformat(v.strip()))
    return max(found, default=None)


def _rows(rows: Any, notes: set[str], where: str, budget: list[int]) -> list[Line]:
    """Spreadsheet rows (values, formulas). The header is the first row with two or more filled cells that are
    all text (datakit/extract.py's rule, Plan 1B Ruling 6); rows above it are plain lines and a stated
    'As of' date there dates every record of the sheet; without one a record is dated by its latest date.
    `budget` is [rows, cells] the whole file may still read, empty rows included: a sheet can name row 10**9
    in a few bytes, and a row naming column XFD is padded to 16,384 cells."""
    out: list[Line] = []
    header: list[Any] | None = None
    stated: date | None = None
    for values_row, formulas_row in rows:
        values = list(values_row)
        budget[0] -= 1
        budget[1] -= max(1, len(values))
        if budget[0] < 0:
            raise IngestError(f"This file has more than {MAX_ROWS:,} rows.")
        if budget[1] < 0:
            raise IngestError(f"This file has more than {MAX_CELLS:,} cells.")
        for i, formula in enumerate(formulas_row):
            if i < len(values) and values[i] is None and isinstance(formula, str) and formula.startswith("="):
                values[i] = FORMULA_NOTE  # adversary F11: never let a computed column vanish silently
                notes.add(f"{where}: formula cells without a saved value; open and save the file in Excel")
        filled = [v for v in values if v not in (None, "")]
        if not filled:
            continue
        if header is None and len(filled) >= 2 and all(isinstance(v, str) for v in filled):
            header = values
            continue
        if header is None:
            line = normalize(" ".join(cell_text(v) for v in filled))
            if m := _AS_OF.search(line):
                with contextlib.suppress(ValueError):  # "As of 2026-13-45" states nothing
                    stated = date.fromisoformat(m.group(1))
            out.append(Line(line))
        else:
            out.append(Line(record_line(header, values), "record", stated or _latest_date(header, values)))
        if len(out) > MAX_LINES:
            raise IngestError(f"This file has more than {MAX_LINES:,} lines.")
    return out


def _xlsx_lines(data: bytes, notes: set[str]) -> list[Line]:
    values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        out: list[Line] = []
        budget = [MAX_ROWS, MAX_CELLS]
        for ws_v, ws_f in zip(values.worksheets, formulas.worksheets, strict=True):
            # The declared size is only a claim: too small hides rows, huge pads empty rows to 16,384 cells.
            ws_v.reset_dimensions()
            ws_f.reset_dimensions()
            rows = zip(ws_v.iter_rows(values_only=True), ws_f.iter_rows(values_only=True), strict=False)
            out += _rows(rows, notes, f"sheet {ws_v.title}", budget)
            if len(out) > MAX_LINES:
                raise IngestError(f"This file has more than {MAX_LINES:,} lines.")
        return [x for x in out if x.text]
    finally:
        values.close()
        formulas.close()


def _csv_lines(text: str, notes: set[str]) -> list[Line]:
    # newline="": a bare CR (old Macintosh files) ends a row too, as csv expects
    rows = ((r, ()) for r in csv.reader(io.StringIO(text, newline="")))
    try:
        return [x for x in _rows(rows, notes, "csv", [MAX_ROWS, MAX_CELLS]) if x.text]
    except csv.Error as exc:  # a cell over the csv module's 131,072-character limit
        raise IngestError(f"This file has a line longer than {MAX_LINE_CHARS:,} characters.") from exc


def sniff(filename: str, data: bytes) -> str:
    """The format by content, which must agree with the extension (spec 9)."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if data.startswith(b"%PDF-"):
        found = "pdf"
    elif data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                infos = z.infolist()
        except (zipfile.BadZipFile, ValueError) as exc:  # ValueError: a member name that is not valid UTF-8
            raise IngestError("This file is damaged and cannot be opened.") from exc
        if len(infos) > MAX_MEMBERS or sum(i.file_size for i in infos) > MAX_UNZIPPED:
            raise IngestError("This file is too large once unpacked.")
        names = {i.filename for i in infos}
        found = "docx" if "word/document.xml" in names else "xlsx" if "xl/workbook.xml" in names else "zip"
        if found == "docx" and any(
            i.file_size > MAX_DOCX_PART
            for i in infos
            if i.filename in ("word/document.xml", "word/styles.xml")
        ):
            raise IngestError("This document is too large to read.")
    elif data.startswith(_OLE2):
        raise IngestError(
            "This looks like an old-format or password-protected Office file. "
            "Save it as .docx or .xlsx without a password and upload it again."
        )
    elif b"\x00" in data:
        found = "binary"
    else:
        found = ext if ext in _TEXT_FORMATS else "text"
    if found != ext or found not in ("pdf", "docx", "xlsx", "csv", "md", "txt"):  # a sentinel never passes
        raise IngestError(
            "Only PDF, DOCX, XLSX, CSV, MD and TXT files are accepted, "
            "and the content must match the extension."
        )
    return found


def decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")  # Windows "Save as CSV" without UTF-8


def parse(filename: str, data: bytes) -> ParsedDocument:
    """Raises IngestError for anything the app will not read."""
    if len(data) > MAX_BYTES:
        raise IngestError("Files must be 4 MB or smaller.")
    fmt = sniff(filename, data)
    notes: set[str] = set()
    try:
        if fmt == "pdf":
            from app.ingest.pdf import pdf_lines  # pdfium loads only for PDFs

            lines = pdf_lines(data)
        elif fmt == "docx":
            lines = _docx_lines(data)
        elif fmt == "xlsx":
            lines = _xlsx_lines(data, notes)
        elif fmt == "csv":
            lines = _csv_lines(decode(data), notes)
        else:
            lines = text_lines(decode(data))
    except IngestError:
        raise
    except Exception as exc:  # the readers raise whatever a damaged package holds: a refusal, not a 500
        raise IngestError("This file is damaged and cannot be opened.") from exc
    if not lines:
        raise IngestError("No text was found in this file.")
    if len(lines) > MAX_LINES:
        raise IngestError(f"This file has more than {MAX_LINES:,} lines.")
    if any(len(x.text) > MAX_LINE_CHARS for x in lines):
        raise IngestError(f"This file has a line longer than {MAX_LINE_CHARS:,} characters.")
    return ParsedDocument(fmt, tuple(lines), tuple(sorted(notes)))  # type: ignore[arg-type]

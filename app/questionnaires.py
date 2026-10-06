"""Questionnaire import (spec 6.10): read an xlsx or csv as text cells, detect the sheet, header row and
columns by header keywords plus content shape, and read the items. The visitor confirms the mapping in a
preview before any item exists. Export lives in app/export.py."""

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from app.api.schemas import Mapping, PreviewRow
from app.contracts import ItemInput
from app.ingest.parse import MAX_BYTES, IngestError, decode, sniff

MAX_ITEMS = 150  # spec 9
MAX_ROWS = 2000  # rows read per sheet (a 150-item questionnaire with section rows fits many times over)
MAX_COLS = 52
MAX_QUESTION = 2000  # characters; the stance and draft prompts carry the question verbatim
MAX_TOPIC = 200
MAX_UNZIPPED = 4 * 1024 * 1024  # export loads the whole workbook, so a bigger one once unpacked is refused
HEADER_SCAN = 30
HEADER_CELL = 40  # a header cell is short; a question is not
PREVIEW = 8
SAMPLE_DIR = Path(__file__).resolve().parent.parent / "data" / "questionnaires"
SAMPLES = {"vsq-a": "vsq-a.xlsx", "mvsp-b": "mvsp-b.csv"}
ROLES: tuple[tuple[str, re.Pattern[str]], ...] = (
    # A whole-cell id header, or one ending in "id" ("Question ID", "Control ID"); never "Yes / No".
    (
        "id",
        re.compile(
            r"^(?:#|id|ref\.?|reference|no\.?|nr\.?|number|code|item"
            r"|n[o\N{MASCULINE ORDINAL INDICATOR}\N{DEGREE SIGN}]\.?)$|\bid$"
        ),
    ),
    ("question", re.compile(r"question|pregunta|frage|requirement|query")),
    ("answer", re.compile(r"answer|response|respuesta|reply|yes ?/ ?no|compliant")),
    ("comments", re.compile(r"comment|note|detail|remark|explanation|evidence|comentario|observaci")),
    ("topic", re.compile(r"domain|section|category|area|topic|control|dominio|secci")),
)


@dataclass(frozen=True)
class Sheet:
    name: str | None  # None for csv
    rows: list[list[str]]  # text cells, row 1 first, padded to the same width


@dataclass(frozen=True)
class ParsedItem:
    row: int
    code: str | None
    question: str
    topic: str | None
    answer: str | None


def _text(value: object) -> str:
    return "" if value is None else " ".join(str(value).split())


def _fold(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text.casefold()) if not unicodedata.combining(c))


def delimiter(text: str) -> str:
    """The header line's most frequent of , ; tab | (csv.Sniffer misreads questions full of commas)."""
    first = text.split("\n", 1)[0]
    return max(",;\t|", key=first.count)


def _pad(rows: list[list[str]]) -> list[list[str]]:
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def read_sheets(filename: str, data: bytes) -> tuple[str, list[Sheet]]:
    """Raises IngestError (shown as is) for a file that is not a readable xlsx or csv."""
    if len(data) > MAX_BYTES:
        raise IngestError("Files must be 4 MB or smaller.")
    fmt = sniff(filename, data)  # content must match the extension; zip size and member caps
    if fmt not in ("xlsx", "csv"):
        raise IngestError("A questionnaire must be an xlsx or csv file.")
    try:
        if fmt == "csv":
            text = decode(data)
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter(text))
            raw = [r for _, r in zip(range(MAX_ROWS + 1), reader, strict=False)]
            if len(raw) > MAX_ROWS or any(len(r) > MAX_COLS for r in raw):  # what export would refuse
                raise IngestError(
                    f"A csv questionnaire can have at most {MAX_ROWS} rows and {MAX_COLS} columns."
                )
            return fmt, [Sheet(None, _pad([[c.strip() for c in r] for r in raw]))]
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                raise IngestError("This workbook is too large to write answers back into.")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = []
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    continue  # hidden sheets are not the visitor's questionnaire (hidden rows are kept)
                ws.reset_dimensions()
                rows = [
                    [_text(v) for v in r[:MAX_COLS]]
                    for _, r in zip(range(MAX_ROWS), ws.iter_rows(values_only=True), strict=False)
                ]
                sheets.append(Sheet(ws.title, _pad(rows)))
            return fmt, sheets
        finally:
            wb.close()
    except IngestError:
        raise
    except Exception as exc:  # a damaged package: a refusal, not a 500 (as parse())
        raise IngestError("This file is damaged and cannot be opened.") from exc


def _role(cell: str) -> str | None:
    if not cell or len(cell) > HEADER_CELL:
        return None
    folded = _fold(cell)
    return next((name for name, rx in ROLES if rx.search(folded)), None)


def _header_like(row: list[str]) -> bool:
    filled = [c for c in row if c]
    return bool(filled) and all(_role(c) for c in filled)


def _columns(header: list[str], below: list[str] | None) -> tuple[dict[str, int], bool]:
    """Role -> 0-based column, first match wins per role; True when the row below joined the header."""
    two_rows = below is not None and _header_like(below)
    roles: dict[str, int] = {}
    for i, cell in enumerate(header):
        joined = f"{cell} {below[i]}".strip() if two_rows and below is not None else cell
        role = _role(cell) or (_role(below[i]) if two_rows and below is not None else None) or _role(joined)
        if role and role not in roles:
            roles[role] = i
    if two_rows and below is not None and "question" in roles and below[roles["question"]]:
        return _columns(header, None)  # the next row holds a question: it is data, not a header
    return roles, two_rows


def detect(sheets: list[Sheet], fmt: str) -> Mapping | None:
    best: tuple[int, int, int, Sheet, dict[str, int]] | None = None  # (-roles, sheet index, row, ...)
    for si, sheet in enumerate(sheets):
        for ri, row in enumerate(sheet.rows[:HEADER_SCAN]):
            below = sheet.rows[ri + 1] if ri + 1 < len(sheet.rows) else None
            roles, _ = _columns(row, below)
            if "question" not in roles:
                continue
            key = (-len(roles), si, ri)
            if best is None or key < best[:3]:
                best = (*key, sheet, roles)
    if best is None:
        return _by_shape(sheets)
    _, _, ri, sheet, roles = best

    def col(role: str) -> str | None:
        return get_column_letter(roles[role] + 1) if role in roles else None

    if "answer" not in roles:  # the first empty-headed column right of the question
        q = roles["question"]
        roles["answer"] = next((i for i in range(q + 1, len(sheet.rows[ri])) if not sheet.rows[ri][i]), q + 1)
    return Mapping(
        sheet=sheet.name,
        header_row=ri + 1,
        id_col=col("id"),
        question_col=col("question") or "A",
        answer_col=col("answer") or "B",
        comments_col=col("comments"),
        topic_col=col("topic"),
    )


def _by_shape(sheets: list[Sheet]) -> Mapping | None:
    """No question header: the column with the most cells ending in "?" (content shape, spec 6.10)."""
    for sheet in sheets:
        width = len(sheet.rows[0]) if sheet.rows else 0
        counts = [sum(r[c].endswith("?") for r in sheet.rows[1:]) for c in range(width)]
        if not counts or max(counts) < 3:
            continue
        q = counts.index(max(counts))
        first = next(i for i, r in enumerate(sheet.rows) if r[q].endswith("?"))
        return Mapping(
            sheet=sheet.name,
            header_row=max(1, first),
            id_col=None,
            question_col=get_column_letter(q + 1),
            answer_col=get_column_letter(q + 2),
            comments_col=None,
        )
    return None


def _sheet(sheets: list[Sheet], mapping: Mapping) -> Sheet:
    for s in sheets:
        if s.name == mapping.sheet:
            return s
    raise IngestError("That sheet is not in the file.")


def read_items(sheets: list[Sheet], mapping: Mapping) -> list[ParsedItem]:
    sheet = _sheet(sheets, mapping)
    h = mapping.header_row - 1
    below = sheet.rows[h + 1] if h + 1 < len(sheet.rows) else None
    _, two_rows = _columns(sheet.rows[h], below) if h < len(sheet.rows) else ({}, False)
    start = h + (2 if two_rows else 1)

    def at(row: list[str], letter: str | None) -> str | None:
        if letter is None:
            return None
        i = column_index_from_string(letter) - 1
        return (row[i] or None) if i < len(row) else None

    items: list[ParsedItem] = []
    section: str | None = None
    for n, row in enumerate(sheet.rows[start:], start=start + 1):
        question = at(row, mapping.question_col)
        if not question:
            filled = [c for c in row if c]
            if filled:
                section = filled[0][:MAX_TOPIC]  # a label the visitor never mapped is cut, not refused
            continue
        topic = at(row, mapping.topic_col) if mapping.topic_col else section
        if len(question) > MAX_QUESTION:
            raise IngestError(f"Row {n}'s question is longer than {MAX_QUESTION:,} characters.")
        if topic and len(topic) > MAX_TOPIC:
            raise IngestError(f"Row {n}'s topic is longer than {MAX_TOPIC} characters.")
        items.append(
            ParsedItem(
                n, at(row, mapping.id_col), " ".join(question.split()), topic, at(row, mapping.answer_col)
            )
        )
        if len(items) > MAX_ITEMS:
            raise IngestError(f"A questionnaire can have at most {MAX_ITEMS} questions.")
    return items


def preview(sheets: list[Sheet], mapping: Mapping) -> list[PreviewRow]:
    """Raises read_items' IngestError, so an upload is refused at once, not at the mapping's PUT."""
    items = read_items(sheets, mapping)
    return [
        PreviewRow(row=i.row, id=i.code, question=i.question, topic=i.topic, answer=i.answer)
        for i in items[:PREVIEW]
    ]


def item_inputs(items: list[ParsedItem]) -> list[ItemInput]:
    """What the engine sees (ItemInput keyed by code); the parity test compares it with the eval's items."""
    return [ItemInput(i.code or str(i.row), i.question, i.topic) for i in items]

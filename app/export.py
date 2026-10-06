"""Export (spec 6.10): write the answers back into the visitor's own file. openpyxl keeps styles, merged
cells, column widths, data validation and conditional formatting; it drops embedded images and charts (a known
openpyxl limit, stated in NOTICE). csv in, csv out. Every cell written is inert text (adversary 1 I1)."""

import csv
import io
import re
import uuid
from copy import copy
from dataclasses import dataclass
from typing import Any

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.schemas import Mapping
from app.db.models import Answer, Item
from app.ingest.parse import decode
from app.questionnaires import delimiter

LABEL_WORDS = {
    "verified": "Verified",
    "partial": "Partial",
    "conflict": "Conflict",
    "unknown": "Unknown",
    "user_confirmed": "Confirmed by you",
    "na": "Not applicable",
}
DRAFT = "Draft, not approved"
APPROVED = "Approved"
ADDED = ("Status", "Sources", "Notes")
NOTICE = "Embedded images and charts are not kept in the exported workbook (an openpyxl limit)."
_ROW = re.compile(r"(\d+)$")
_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")


@dataclass(frozen=True)
class ExportRow:
    row: int
    label: str
    value: str | None
    text: str
    approved: bool
    sources: list[str]
    notes: list[str]


def rows_for(session: Session, run_id: uuid.UUID) -> list[ExportRow]:
    out = []
    pairs = session.execute(
        select(Item, Answer)
        .join(Answer, Answer.item_id == Item.id)
        .where(Answer.run_id == run_id)
        .order_by(Item.position)
    )
    for item, a in pairs:
        m = _ROW.search(item.row_ref)
        if m is None:
            continue
        sources = list(dict.fromkeys(f"{c['filename']} line {c['line_start']}" for c in a.citations))
        notes = [n for n in (a.scope_note,) if n]
        value = "N/A" if a.label == "na" else a.value
        out.append(
            ExportRow(int(m.group(1)), a.label, value, a.text, a.approved_at is not None, sources, notes)
        )
    return out


def _status(r: ExportRow) -> str:
    return f"{LABEL_WORDS[r.label]} · {APPROVED if r.approved else DRAFT}"


def _csv_safe(value: str) -> str:
    return "'" + value if value.startswith(_TRIGGERS) else value


def _put(cell: Any, value: str | None) -> None:
    """Write a string that stays a string: openpyxl would turn a leading = into a formula (data_type 'f')."""
    cell.value = value
    if value is not None:
        cell.data_type = "s"


def _cells(
    r: ExportRow, allowed: set[str] | None, has_comments: bool, comments_empty: bool
) -> tuple[str | None, str | None, list[str]]:
    """(answer cell, comments cell, notes) for one row."""
    notes = list(r.notes)
    answer = r.value or (r.text or None)
    if answer is not None and allowed is not None and answer not in allowed:
        # A validation list (Yes/No/N/A) refuses the value or the text: never write it into that cell.
        if r.value:
            notes.insert(0, f"Answer: {r.value}")
        answer = None
    put_text = bool(r.text) and answer != r.text
    to_comments = put_text and has_comments and comments_empty
    if put_text and not to_comments:
        notes.append(f"Answer text: {r.text}")
    return answer, (r.text if to_comments else None), notes


def _validation(ws: Any, letter: str) -> set[str] | None:
    for dv in ws.data_validations.dataValidation:
        if (
            dv.type == "list"
            and dv.formula1
            and any(letter == re.sub(r"\d", "", str(c).split(":")[0]) for c in str(dv.sqref).split())
        ):
            return {v.strip() for v in dv.formula1.strip('"').split(",")}
    return None


def export_xlsx(original: bytes, mapping: Mapping, rows: list[ExportRow]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(original))
    ws = wb[mapping.sheet] if mapping.sheet else wb.active
    h = mapping.header_row
    first = ws.max_column + 1
    style = ws.cell(h, column_index_from_string(mapping.question_col))
    for i, title in enumerate(ADDED):
        cell = ws.cell(h, first + i, title)
        cell.font, cell.fill, cell.border, cell.alignment = (
            copy(style.font),
            copy(style.fill),
            copy(style.border),
            copy(style.alignment),
        )
    allowed = _validation(ws, mapping.answer_col)
    for r in rows:
        comments = f"{mapping.comments_col}{r.row}" if mapping.comments_col else None
        answer, comment, notes = _cells(
            r, allowed, comments is not None, comments is None or ws[comments].value in (None, "")
        )
        _put(ws[f"{mapping.answer_col}{r.row}"], answer)
        if comments and comment:
            _put(ws[comments], comment)
        _put(ws.cell(r.row, first), _status(r))
        _put(ws.cell(r.row, first + 1), "; ".join(r.sources) or None)
        _put(ws.cell(r.row, first + 2), " ".join(notes) or None)
    for i in range(3):
        ws.column_dimensions[get_column_letter(first + i)].width = 28
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_csv(original: bytes, mapping: Mapping, rows: list[ExportRow]) -> bytes:
    text = decode(original)
    sep = delimiter(text)
    table = list(csv.reader(io.StringIO(text, newline=""), delimiter=sep))
    width = max(len(r) for r in table)
    table = [r + [""] * (width - len(r)) for r in table]
    a = column_index_from_string(mapping.answer_col) - 1
    c = column_index_from_string(mapping.comments_col) - 1 if mapping.comments_col else None
    table[mapping.header_row - 1] += list(ADDED)
    by_row = {r.row: r for r in rows}
    for n, line in enumerate(table, start=1):
        if n == mapping.header_row:
            continue
        r = by_row.get(n)
        if r is None:
            line += ["", "", ""]
            continue
        answer, comment, notes = _cells(r, None, c is not None, c is None or not line[c])
        line[a] = _csv_safe(answer or "")
        if c is not None and comment:
            line[c] = _csv_safe(comment)
        line += [_csv_safe(s) for s in (_status(r), "; ".join(r.sources), " ".join(notes))]
    out = io.StringIO(newline="")
    csv.writer(out, delimiter=sep, lineterminator="\r\n").writerows(table)
    bom = "\N{ZERO WIDTH NO-BREAK SPACE}" if original.startswith(b"\xef\xbb\xbf") else ""
    return (bom + out.getvalue()).encode("utf-8")

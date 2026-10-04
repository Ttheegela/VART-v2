from datetime import date
from itertools import pairwise
from pathlib import Path

import docx
import openpyxl
from fpdf import FPDF

from datakit.extract import lines_of, text_of


def test_docx_paragraphs_and_tables_in_order(tmp_path: Path) -> None:
    d = docx.Document()
    d.add_heading("Access Control", level=1)
    d.add_paragraph("User access is reviewed quarterly.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "System", "Owner"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Okta", "Dana Ortiz"
    d.add_paragraph("After the table.")
    path = tmp_path / "a.docx"
    d.save(str(path))
    assert lines_of(path) == [
        "Access Control",
        "User access is reviewed quarterly.",
        "System: Okta; Owner: Dana Ortiz",
        "After the table.",
    ]


def test_xlsx_rows_become_record_lines_after_the_header(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Zed access review log"])
    ws.append(["As of: 2026-09-15"])
    ws.append([])
    ws.append(["System", "Last review completed", "Status"])
    ws.append(["Okta", "2026-01-10", "Overdue"])
    path = tmp_path / "a.xlsx"
    wb.save(str(path))
    assert lines_of(path) == [
        "Zed access review log",
        "As of: 2026-09-15",
        "System: Okta; Last review completed: 2026-01-10; Status: Overdue",
    ]


def test_pdf_wrapped_lines_still_contain_whole_sentences(tmp_path: Path) -> None:
    words = [f"word{i:03d}" for i in range(120)]  # distinct, so a split, a lost or a repeated word shows
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 5, " ".join(words), new_x="LMARGIN", new_y="NEXT")
    path = tmp_path / "a.pdf"
    pdf.output(str(path))
    assert lines_of(path) == [" ".join(words)]  # wrapping is not a paragraph break
    text = text_of(path)
    assert all(f"{a} {b}" in text for a, b in pairwise(words))


def test_xlsx_rows_above_the_header_read_like_record_cells(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Zed log", date(2026, 9, 15), 42.0])
    ws.append(["System", "Status"])
    ws.append(["Okta", "Overdue"])
    path = tmp_path / "a.xlsx"
    wb.save(str(path))
    assert lines_of(path) == ["Zed log 2026-09-15 42", "System: Okta; Status: Overdue"]


def test_markdown_tables_and_bullets(tmp_path: Path) -> None:
    path = tmp_path / "a.md"
    path.write_text("# Title\n\nOne para\ncontinues here.\n\n- A bullet\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    assert lines_of(path) == ["Title", "One para continues here.", "A bullet", "A: 1; B: 2"]

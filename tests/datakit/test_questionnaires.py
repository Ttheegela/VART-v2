import csv
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.packaging import extended

from datakit import validate
from datakit.questionnaires import build_csv, build_xlsx, mvsp_items, vsaq_items
from datakit.schemas import Facts, Selection, load_yaml
from datakit.validate import DATA, STAGES

Q = DATA / "questionnaires"


def _sel(name: str) -> Selection:
    return load_yaml(Q / f"{name}.selection.yaml", Selection)


def test_sources_parse() -> None:
    assert len(vsaq_items()) >= 150  # 199 question items with an id and text at the pinned commit
    assert len(mvsp_items()) == 25


def test_every_selected_item_has_a_real_source_and_a_known_control() -> None:
    known = vsaq_items() | mvsp_items()
    controls = {c.id for c in load_yaml(DATA / "dev" / "facts.yaml", Facts).controls}
    for name in ("vsq-a", "mvsp-b"):
        for it in _sel(name).items:
            assert it.source in known, it.code
            assert it.control in controls, it.code


def test_sizes_and_unique_codes() -> None:
    a, b = _sel("vsq-a"), _sel("mvsp-b")
    assert 55 <= len(a.items) <= 65 and len(b.items) == 25
    for sel in (a, b):
        codes = [i.code for i in sel.items]
        assert len(codes) == len(set(codes))


def test_xlsx_has_the_messy_layout(tmp_path: Path) -> None:
    path = tmp_path / "a.xlsx"
    build_xlsx(_sel("vsq-a"), path)
    wb = openpyxl.load_workbook(str(path))
    assert wb.sheetnames == ["Instructions", "Questionnaire"]
    ws = wb["Questionnaire"]
    assert [ws.cell(5, c).value for c in range(1, 6)] == [
        "#",
        "Domain",
        "Control Question",
        "Response (Yes / No / N/A)",
        "Comments / Evidence",
    ]
    assert ws.merged_cells.ranges, "section rows are merged across the table"
    assert ws.data_validations.dataValidation, "the response column has a Yes/No/N/A list"
    assert ws.freeze_panes == "A6"


def test_builds_are_byte_identical(tmp_path: Path) -> None:
    build_xlsx(_sel("vsq-a"), tmp_path / "1.xlsx")
    build_xlsx(_sel("vsq-a"), tmp_path / "2.xlsx")
    digest = [hashlib.sha256((tmp_path / n).read_bytes()).hexdigest() for n in ("1.xlsx", "2.xlsx")]
    assert digest[0] == digest[1]


def test_csv_round_trips(tmp_path: Path) -> None:
    path = tmp_path / "b.csv"
    build_csv(_sel("mvsp-b"), path)
    with path.open(newline="") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == ["ID", "Control", "Question", "Answer", "Notes"] and len(rows) == 25


# The tests below pin the selection rules from the plan that no other check enforces.
SECTIONS = {
    "Governance",
    "Access Control",
    "Data Security",
    "Business Continuity",
    "Vulnerability Management",
    "Secure Development",
    "Incident Response",
    "Vendor Management",
    "Human Resources",
    "Logging and Monitoring",
    "Asset Management",
    "Engagement",
}
# Every wording of a conflict-trap control carries the threshold the documents disagree on.
THRESHOLDS = {
    "access-review": ["at least quarterly"],
    "backup-restore-test": ["at least quarterly"],
    "log-retention": ["at least 12 months"],
    "cloud-only": ["ALL infrastructure", "AWS", "on premises"],
    "data-residency": ["ALL customer data", "AWS", "on premises"],
}


def _asked(name: str) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for it in _sel(name).items:
        out.setdefault(it.control, []).append(it.question)
    return out


def test_vsq_a_covers_every_control_and_words_the_trap_controls_with_their_thresholds() -> None:
    controls = {c.id for c in load_yaml(DATA / "dev" / "facts.yaml", Facts).controls}
    a = _asked("vsq-a")
    assert set(a) == controls
    assert all(len(a[c]) == 2 for c in ("mfa", "access-review", "backup-restore-test"))
    for asked in (a, _asked("mvsp-b")):
        for control, phrases in THRESHOLDS.items():
            for question in asked.get(control, []):
                assert all(p in question for p in phrases), (control, question)


def test_mvsp_b_is_one_item_per_mvsp_control() -> None:
    items = _sel("mvsp-b").items
    assert [i.source for i in items] == list(mvsp_items())
    assert all(i.code == i.source.replace("mvsp:", "MVSP-") for i in items)


def test_sections_csf_ids_and_ascii() -> None:
    assert {i.section for i in _sel("vsq-a").items} <= SECTIONS
    csf = re.compile(r"(GV|ID|PR|DE|RS|RC)\.[A-Z]{2}-\d{2}")
    for name in ("vsq-a", "mvsp-b"):
        assert (Q / f"{name}.selection.yaml").read_bytes().isascii()
        for it in _sel(name).items:
            assert (it.csf_id is None) == (it.section == "Engagement"), it.code  # no CSF ID for engagement
            assert it.csf_id is None or csf.fullmatch(it.csf_id), it.code


def test_committed_files_are_current_and_the_mapping_matches_the_layout() -> None:
    assert STAGES["questionnaires"]("dev") == []
    mapping = json.loads((Q / "vsq-a.mapping.json").read_text(encoding="utf-8"))
    ws = openpyxl.load_workbook(str(Q / "vsq-a.xlsx"))[mapping["sheet"]]
    columns = [mapping[k] for k in ("id_col", "question_col", "answer_col", "comments_col")]
    header = [ws[f"{c}{mapping['header_row']}"].value for c in columns]
    assert header == ["#", "Control Question", "Response (Yes / No / N/A)", "Comments / Evidence"]
    assert (Q / "mvsp-b.csv").read_bytes().isascii()


def test_xlsx_bytes_do_not_depend_on_the_openpyxl_version_or_on_zlib(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    build_xlsx(_sel("vsq-a"), tmp_path / "now.xlsx")
    monkeypatch.setattr(extended, "__version__", "9.9.9")
    openpyxl.Workbook().save(str(tmp_path / "raw.xlsx"))
    with zipfile.ZipFile(tmp_path / "raw.xlsx") as z:
        # the patch reaches openpyxl's writer, so the comparison below means something
        assert b"9.9.9" in z.read("docProps/app.xml")
    build_xlsx(_sel("vsq-a"), tmp_path / "later.xlsx")
    assert (tmp_path / "now.xlsx").read_bytes() == (tmp_path / "later.xlsx").read_bytes()
    with zipfile.ZipFile(tmp_path / "now.xlsx") as z:
        # stored members: there is no compressor left to differ between zlib builds
        assert {i.compress_type for i in z.infolist()} == {zipfile.ZIP_STORED}
        app = z.read("docProps/app.xml")
    assert b"<Application>" not in app and b"<AppVersion>" not in app


def test_the_stage_flags_a_stale_mapping(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    copy = tmp_path / "questionnaires"
    shutil.copytree(Q, copy)
    monkeypatch.setattr(validate, "QDIR", copy)
    assert STAGES["questionnaires"]("dev") == []
    mapping = copy / "vsq-a.mapping.json"
    mapping.write_text(mapping.read_text(encoding="utf-8").replace('"header_row": 5', '"header_row": 6'))
    assert STAGES["questionnaires"]("dev") == [
        "vsq-a.mapping.json is stale: run python -m datakit.questionnaires"
    ]


# Controls that were must-ask and are now documented: each has one plain yes sentence (no negation cue words).
DOCUMENTED = {
    "vuln-disclosure": "faq",
    "media-sanitization": "amp",
    "build-release": "sdp",
    "secure-frameworks": "sdp",
}
CUES = re.compile(r"\b(not|no|never|none|without|pending|planned)\b", re.IGNORECASE)


def test_documented_controls_have_one_plain_yes_sentence() -> None:
    facts = load_yaml(DATA / "dev" / "facts.yaml", Facts)
    for control, doc in DOCUMENTED.items():
        (s,) = facts.statements_for(control)
        assert (s.doc, s.stance, s.flags) == (doc, "yes", ()), control
        assert s.text.isascii() and s.text.endswith(".") and not CUES.search(s.text), s.text
        assert facts.control(control).truth != "(not in any document)", control
    assert not [t.id for t in facts.traps if t.kind == "must_ask" and set(t.controls) & set(DOCUMENTED)]


# MVSP-B questions that must stay faithful to their MVSP control: (control, must keep, must not have).
MVSP_B = {
    "MVSP-1.4": ("pentest", ["penetration test"], ["comprehensive", "services"]),
    "MVSP-1.5": ("security-training", ["role-specific"], []),
    "MVSP-2.1": ("customer-sso", ["customers"], ["employees", "internal"]),
    "MVSP-2.6": ("patch-sla", ["dependencies", "patch"], ["scan"]),
    "MVSP-3.3": ("developer-training", ["train your developers"], ["static", "SAST"]),
    "MVSP-3.5": ("build-release", ["CI/CD"], []),
}


def test_mvsp_b_questions_keep_the_mvsp_wording_and_map_to_the_right_control() -> None:
    items = {i.code: i for i in _sel("mvsp-b").items}
    for code, (control, keep, drop) in MVSP_B.items():
        item = items[code]
        assert item.control == control, code
        assert all(p in item.question for p in keep), (code, item.question)
        assert not any(p in item.question for p in drop), (code, item.question)


def test_contact_questions_ask_for_deal_specific_details() -> None:
    detail = ["Northbeam Health's account", "that person's name, email address and a 24/7 phone number"]
    asked = _asked("vsq-a")
    for control, who in (
        ("security-contact", "a security contact"),
        ("incident-contact", "an incident contact"),
    ):
        (question,) = asked[control]
        assert question.startswith(f"Will you designate {who} for "), control
        assert all(p in question for p in detail), control

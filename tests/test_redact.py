import io
import time
from datetime import date

import docx
import openpyxl
import pytest

from app import redact
from app.contracts import Line
from app.ingest.parse import IngestError, parse
from app.patterns import PLACEHOLDER
from app.redact import redact_lines, redact_text


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        (
            "The program is owned by Dana Ortiz, Head of Security.",
            "The program is owned by <PERSON>, Head of Security.",
        ),
        (
            "System: Okta; Owner: Marcus Lee; Status: Overdue",
            "System: Okta; Owner: <PERSON>; Status: Overdue",
        ),
        ("Write to dana.ortiz@kestrelyn.example today.", "Write to <EMAIL> today."),
        ("Call +1 512 555 0142 or (512) 555-0143.", "Call <PHONE> or <PHONE>."),
        ("Our office is at 1200 Congress Avenue, Suite 400.", "Our office is at <ADDRESS>."),
        ("db: postgres://admin:hunter2pass@db.example.com:5432/app", "db: <SECRET>"),
        ("key sk-or-v1-abcdefghijklmnopqrstuvwxyz012345", "key <SECRET>"),
        ("aws AKIAIOSFODNN7EXAMPLE", "aws <SECRET>"),
        ("password = hunter2hunter2", "<SECRET>"),
        ("José Álvarez signed the policy.", "<PERSON> signed the policy."),
        ("Owner: Dana M. Ortiz; Status: Open", "Owner: <PERSON>; Status: Open"),
        ("Signed by DANA ORTIZ", "Signed by <PERSON>"),
        ("AWS_SECRET_ACCESS_KEY=" + "wJalrXUtnFEMI/K7MDENG/bPxRfiCY" + "EXAMPLEKEY", "<SECRET>"),
        ("aws_secret_access_key = " + "wJalrXUtnFEMI/K7MDENG/bPxRfiCY" + "EXAMPLEKEY", "<SECRET>"),
        ("DB_PASSWORD=hunter2hunter2", "<SECRET>"),
        ("API_KEY_PROD=abcdefgh12345678", "<SECRET>"),
        ("API_KEY_LIVE=abcdefgh12345678", "<SECRET>"),
        ("X-Api-Key-Prod=abcdefgh12345678", "<SECRET>"),
        ("SECRET_KEY_BASE=abcdefgh12345678", "<SECRET>"),
        ("CLIENT_SECRET_V2=abcdefgh12345678", "<SECRET>"),
        ("access_token_v2=abcdefgh12345678", "<SECRET>"),
        ("DB_PASSWORD_PROD=abcdefgh12345678", "<SECRET>"),
        ("DB_PASSWORD_FILE=abcdefgh12345678", "<SECRET>"),
        ("PASSWORD_HASH=abcdefgh12345678", "<SECRET>"),
        ("PASSWORD_SALT=abcdefgh12345678", "<SECRET>"),
        ("SECRET_PASSPHRASE=abcdefgh12345678", "<SECRET>"),
        ("SECRET_BASE64=abcdefgh12345678", "<SECRET>"),
        ("secretKey: abcdefgh12345678", "<SECRET>"),
        ("accessToken: abcdefgh12345678", "<SECRET>"),
        ("client_secret: abcdefgh12345678", "<SECRET>"),
        ("pwd: hunter2hunter2", "<SECRET>"),
        ("Authorization: Bearer " + "abcdefghijklmnopqrstuvwxyz012345", "Authorization: <SECRET>"),
        ("key " + "sk_" + "live_" + "abcdefghijklmnopqrstuvwx", "key <SECRET>"),
        ("key " + "sk_" + "test_" + "abcdefghijklmnopqrstuvwx", "key <SECRET>"),
        ("key " + "rk_" + "live_" + "abcdefghijklmnopqrstuvwx", "key <SECRET>"),
        ("token " + "github_" + "pat_" + "abcdefghijklmnopqrstuvwxyz0123", "token <SECRET>"),
        ("token " + "ghp_" + "abcdefghijklmnopqrstuvwxyz0123456789", "token <SECRET>"),
        ("hook " + "https://hooks.slack" + ".com/services/T000/B000/abcdefgh", "hook <SECRET>"),
        ("url?" + "sig=" + "abcdefghijklmnop%2Bqrstu&se=1", "url?<SECRET>&se=1"),
    ],
)
def test_private_data_and_secrets_are_replaced(text: str, expected: str) -> None:
    assert redact_text(text) == expected


@pytest.mark.parametrize(
    "text",
    [
        "Kestrelyn maintains a documented incident response plan.",  # a company name is not a person
        "Kestrelyn Ledger stores all customer data in AWS us-east-1.",  # nor a product name
        "Sablecrest Security performed the test.",
        "Customer data is stored in Ireland.",  # place names stay: data-residency answers need them
        "Version 2.0 was released on 2026-01-10 for ISO 27001:2022.",  # dates and versions are not phones
        "Google Workspace and GitHub are reviewed quarterly.",
        "password complexity: minimum 12 characters",
        "secret = see vault",
        "API key: rotated-every-90-days",
        "INFORMATION SECURITY POLICY",
        "Passwords: bcrypt-hashed with per-user salts",
        "Tokenisation: Format-preserving encryption for card data",
        "Secrets: HashiCorp Vault",
        "Session tokens: HttpOnly-and-Secure",
        "token_count: 1000000000",
        # adversary checkpoint 3, I5: addresses, years and counts are not phones; policy values not secrets
        "Internal range 192.168.100.200 and 172.16.254.100 are segmented.",
        "Allowed egress: 10.100.200.0/24 only.",
        "Years covered: 2024 2025 2026",
        "Budget: 10 000 000 records.",
        "Password: Required; MFA: Yes",
        "Password: bcrypt-hashed, rotated every 90 days.",
        "Secret: HashiCorp Vault, rotated quarterly.",
        "Token: RS256-signed",
        # I1: a hash is not a key body (one case: hex is single-case)
        "e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
        "commit 9298371c0ffee0ddba11ad5eed0f00dfacade123 merged",
        # re-review N1: the extra keyed rule leaves policy cells alone
        "Password: required for all accounts",
        "Password: hashed-and-salted",
        "Password: Required; Secret: Vault; Token: OAuth2",
    ],
)
def test_business_text_is_left_alone(text: str) -> None:
    assert redact_text(text) == text


def test_a_private_key_block_is_one_secret() -> None:
    # Built in pieces so that no file holds a whole fake key block for gitleaks to match (Plan 2 addendum).
    key = "-----BEGIN RSA " + "PRIVATE KEY-----\nMIIEow\nabc\n-----END RSA " + "PRIVATE KEY-----"
    assert redact_text(f"before {key} after") == "before <SECRET> after"


def test_redaction_tokens_never_look_like_template_placeholders() -> None:
    text = redact_text("Owner: Marcus Lee; Contact: lee@example.com; Phone: 512 555 0142")
    assert "<PERSON>" in text and PLACEHOLDER.search(text) is None


def test_lines_keep_their_kind_and_date() -> None:
    (line,) = redact_lines([Line("Owner: Marcus Lee", "record", date(2026, 1, 10))])
    assert (line.text, line.kind, line.as_of) == ("Owner: <PERSON>", "record", date(2026, 1, 10))


def test_overlapping_matches_collapse_to_one_token() -> None:
    # the URL's user:password@host also looks like an email; only the outer span may be replaced
    text = "db: postgres://admin:hunter2pass@db.example.com/app done"
    assert redact_text(text) == "db: <SECRET> done"


def test_digits_glued_to_a_word_are_not_a_phone() -> None:
    assert redact_text("order A512 555 0142 shipped") == "order A512 555 0142 shipped"


@pytest.mark.parametrize(
    "line",
    ["token_" * 3333, "api-key-" * 2500, "password_" * 2200, "secret=" * 2800, "a" * 20000, "A1 " * 6666],
    ids=["token_", "api-key-", "password_", "secret=", "a", "A1 "],
)
def test_a_crafted_line_is_redacted_in_linear_time(line: str) -> None:
    started = time.monotonic()
    redact_text(line)
    # generous: about 1.5 s of this is spaCy on 20,000 characters; the keyed-secret regex was cubic (minutes)
    assert time.monotonic() - started < 10.0


# Built in pieces so that no file holds a whole fake key block for gitleaks to match (Plan 2 addendum).
_PEM = [
    "-----BEGIN RSA " + "PRIVATE KEY-----",
    *["MIIEowIBAAKCAQEA" + "x" * 48] * 3,
    "-----END RSA " + "PRIVATE KEY-----",
]


def _docx_key() -> bytes:
    d = docx.Document()
    d.add_paragraph("Deployment notes")
    for row in _PEM:
        d.add_paragraph(row)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def _xlsx_key() -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(["Name", "Value"])
    sheet.append(["deploy key", _PEM[0]])
    for row in _PEM[1:]:
        sheet.append(["", row])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    ("name", "data"), [("notes.docx", _docx_key()), ("keys.xlsx", _xlsx_key())], ids=["docx", "xlsx"]
)
def test_a_private_key_body_split_over_lines_is_redacted(name: str, data: bytes) -> None:
    # adversary checkpoint 3, I1: docx and xlsx give one line per key row, so the BEGIN line alone matched
    lines = redact_lines(parse(name, data).lines)
    assert not any("MIIE" in line.text for line in lines)
    assert sum("<SECRET>" in line.text for line in lines) >= 4


def test_redaction_past_its_time_budget_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    # adversary checkpoint 3, I6: a crafted upload cost ~590 s of spaCy; a deadline between slices stops it
    monkeypatch.setattr(redact, "DEADLINE_S", -1.0)
    with pytest.raises(IngestError, match="too long to process"):
        redact_lines([Line("Owner: Marcus Lee")] * 20)


def test_lines_redacted_in_slices_match_one_by_one() -> None:
    texts = [f"Row {i}: owner Marcus Lee, phone 512 555 01{i:02d}." for i in range(20)]
    assert [x.text for x in redact_lines([Line(t) for t in texts])] == [redact_text(t) for t in texts]


@pytest.mark.parametrize(
    "text",
    [  # adversary-3 re-review N1: realistic keyed secrets 65ba93c let through (in pieces for gitleaks)
        "DB_PASSWORD=" + "supersecretpassword",
        "POSTGRES_PASSWORD: " + "correcthorsebattery",
        "api_key=" + "1234567890123456,",
        "API_KEY=" + "12345678901234567890",
        "password: " + "12345678.",
        "Password: " + "12345678",
        "token: " + "9876543210987654",
        "password: " + "correcthorsebatterystaple",
        "password: " + "correct-horse-battery-staple",
        "client_secret: " + "correct-horse-battery-staple",
        "password: " + '"correct horse battery staple"',
        "password = " + "'letmein-please-now'",
        "spring.datasource.password=" + "changeme",
        "clientSecret: " + "abcdefghijklmnop",
        "secret=" + "abcdefghijklmnopqrstuvwx",
        "password=" + "Summer-Breeze-2024",
        "DB_PASSWORD=" + "hunter2hunter2",
        "System: Prod DB; Username: admin; Password: " + "correcthorsebattery; Notes: rotate",
        "Password: " + "Tr0ub4dor&3",
        "  password: " + "changemeplease",
        "x-api-key: " + "1234567890123456",
        "export GITHUB_TOKEN=" + "abcdefghijklmnopqrstuvwxyz",
    ],
)
def test_a_keyed_secret_of_any_shape_is_redacted(text: str) -> None:
    out = redact_text(text)
    assert "<SECRET>" in out
    for value in (
        "supersecret",
        "correcthorse",
        "12345678",
        "9876543210",
        "correct-horse",
        "letmein",
        "changeme",
        "abcdefghijklmnop",
        "Summer-Breeze",
        "hunter2",
        "Tr0ub4dor",
    ):
        assert value not in out


@pytest.mark.parametrize(
    "text",
    [
        "PRIVATE_KEY=" + "MIIEowIBAAKCAQEA" + "x" * 48,
        "PRIVATE_KEY_B64=" + "LS0tLS1CRUdJTiBSU0Eg" + "UHJpdmF0ZQ" * 4,
    ],
)
def test_a_key_body_right_after_an_equals_sign_is_redacted(text: str) -> None:
    # adversary-3 re-review N2: the PEM-row lookbehind excluded '='
    assert redact_text(text).endswith("=<SECRET>")

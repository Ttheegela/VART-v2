import io
import zipfile
from datetime import date, datetime
from pathlib import Path

import docx
import openpyxl
import pytest
from docx.oxml import parse_xml

import app.ingest.parse as parse_module
from app.ingest.parse import (
    FORMULA_NOTE,
    MAX_BYTES,
    MAX_LINE_CHARS,
    MAX_LINES,
    IngestError,
    parse,
    text_lines,
)
from datakit.extract import lines_of
from datakit.schemas import Facts, load_yaml

ROOT = Path(__file__).resolve().parent.parent
NL = chr(10)
FACTS = load_yaml(ROOT / "data" / "dev" / "facts.yaml", Facts)


@pytest.mark.parametrize("spec", [d for d in FACTS.documents if d.format != "pdf"], ids=lambda d: d.id)
def test_every_dev_document_reads_like_the_reference_extractor(spec) -> None:  # type: ignore[no-untyped-def]
    path = ROOT / "data" / "dev" / "docs" / spec.filename
    assert [line.text for line in parse(spec.filename, path.read_bytes()).lines] == lines_of(path)


def test_headings_bullets_and_table_rows_get_their_kind() -> None:
    md = "# Access" + NL * 2 + "Users are" + NL + "reviewed quarterly." + NL * 2 + "- MFA is on" + NL * 2
    md += "| System | Owner |" + NL + "|---|---|" + NL + "| Okta | IT |" + NL
    got = [(x.text, x.kind) for x in parse("p.md", md.encode()).lines]
    assert got == [
        ("Access", "heading"),
        ("Users are reviewed quarterly.", "text"),
        ("MFA is on", "text"),
        ("System: Okta; Owner: IT", "record"),
    ]


def _xlsx(*rows: list[object]) -> bytes:
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_stated_as_of_date_wins_over_the_row_dates() -> None:
    data = _xlsx(
        ["Access review log"],
        ["As of: 2026-09-15"],
        [],
        ["System", "Last review"],
        ["Okta", datetime(2026, 1, 10)],
    )
    lines = parse("log.xlsx", data).lines
    assert [(x.text, x.kind, x.as_of) for x in lines] == [
        ("Access review log", "text", None),
        ("As of: 2026-09-15", "text", None),
        ("System: Okta; Last review: 2026-01-10", "record", date(2026, 9, 15)),
    ]


def test_without_a_stated_date_a_row_is_dated_by_its_latest_date() -> None:
    lines = parse(
        "log.xlsx", _xlsx(["System", "Reviewed", "Approved"], ["Okta", "2026-01-10", datetime(2026, 4, 10)])
    ).lines
    assert lines[0].as_of == date(2026, 4, 10)


def test_a_date_still_to_come_never_dates_a_row() -> None:
    data = "System,Last review,Next review" + NL + "Okta,2026-01-10,2027-01-10" + NL
    assert parse("reviews.csv", data.encode()).lines[0].as_of == date(2026, 1, 10)


def test_formula_cells_without_a_saved_value_are_shown_and_noted() -> None:
    parsed = parse("f.xlsx", _xlsx(["System", "Days overdue"], ["Okta", "=1+1"]))
    assert parsed.lines[0].text == f"System: Okta; Days overdue: {FORMULA_NOTE}"
    assert parsed.notes and "formula" in parsed.notes[0]


def test_csv_in_windows_encoding_or_with_a_byte_order_mark() -> None:
    cafe = "caf\N{LATIN SMALL LETTER E WITH ACUTE}"
    windows = f"System,Note\r\nOkta,{cafe} \N{RIGHT SINGLE QUOTATION MARK}ok\r\n".encode("cp1252")
    assert parse("w.csv", windows).lines[0].text == f"System: Okta; Note: {cafe} 'ok"
    bom = "\N{ZERO WIDTH NO-BREAK SPACE}System,Status\r\nOkta,Overdue\r\n".encode()
    assert parse("b.csv", bom).lines[0].text == "System: Okta; Status: Overdue"


def test_a_table_row_in_any_document_is_dated_by_its_latest_date() -> None:
    # Spec 6.4: a table row is a record, dated by its latest date cell; 2026-13-45 is not a date.
    md = NL.join(
        ["| System | Last review | Checked |", "|---|---|---|", "| Okta | 2026-01-10 | 2026-13-45 |", ""]
    )
    (row,) = parse("t.md", md.encode()).lines
    assert (row.text, row.kind, row.as_of) == (
        "System: Okta; Last review: 2026-01-10; Checked: 2026-13-45",
        "record",
        date(2026, 1, 10),
    )


@pytest.mark.parametrize(
    ("item", "line"),
    [("- [x] MFA is enforced", "MFA is enforced"), ("- [ ] MFA is enforced", "[ ] MFA is enforced")],
)
def test_a_checked_task_box_is_dropped_and_an_open_one_kept(item: str, line: str) -> None:
    # Plan 2A Ruling 10: a done item is a plain statement; "[ ]" stays for PLACEHOLDER to flag.
    assert [x.text for x in text_lines(item)] == [line]


def _bomb() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", b"0" * (60 * 1024 * 1024))
    return buf.getvalue()


_OLE2 = bytes.fromhex("D0CF11E0A1B11AE1")  # legacy .doc/.xls and password-protected Office files


def _bad_member_name(real_member: str) -> bytes:
    """A zip with a member whose name bytes (ff fe) are not UTF-8 although the UTF-8 flag (0x800) is set."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(real_member, "<root/>")
        z.writestr("xx", "x")
    data = bytearray(buf.getvalue())
    for header, name_at, flags_at in ((b"PK\x03\x04", 30, 6), (b"PK\x01\x02", 46, 8)):
        start = data.rindex(header)  # the last header of each kind belongs to the member "xx"
        assert data[start + name_at : start + name_at + 2] == b"xx"
        data[start + flags_at + 1] |= 0x08  # flag bit 11
        data[start + name_at : start + name_at + 2] = b"\xff\xfe"
    return bytes(data)


@pytest.mark.parametrize(
    ("name", "data", "message"),
    [
        ("big.txt", b"x" * (MAX_BYTES + 1), "4 MB"),
        ("a.csv", b"a,b\x00c", "content must match"),
        ("a.txt", b"%PDF-1.4 not really", "content must match"),
        ("a.exe", b"MZ\x90\x00", "content must match"),
        ("a.docx", b"PK\x03\x04garbage", "damaged"),
        ("a.xlsx", _bomb(), "too large once unpacked"),
        ("long.txt", b"line.\n\n" * (MAX_LINES + 1), "20,000 lines"),
        ("wide.txt", b"x" * (MAX_LINE_CHARS + 1), "20,000 characters"),
        ("empty.md", b"\n\n   \n", "No text"),
        ("x.text", b"MFA is enforced.", "content must match"),
        ("x.binary", b"MFA\x00is enforced.", "content must match"),
        ("a.docx", _bad_member_name("word/document.xml"), "damaged"),
        ("a.xlsx", _bad_member_name("xl/workbook.xml"), "damaged"),
        ("a.doc", _OLE2 + b"\x00" * 64, "old-format or password-protected"),
        ("a.xlsx", _OLE2 + b"\x00" * 64, "old-format or password-protected"),
    ],
)
def test_files_the_app_will_not_read_get_a_readable_reason(name: str, data: bytes, message: str) -> None:
    with pytest.raises(IngestError, match=message):
        parse(name, data)


def test_a_docx_renamed_to_xlsx_is_refused() -> None:
    docx_bytes = (ROOT / "data" / "dev" / "docs" / "access-control-policy.docx").read_bytes()
    with pytest.raises(IngestError, match="content must match"):
        parse("policy.xlsx", docx_bytes)


def test_a_zip_archive_is_refused_even_when_its_name_says_zip() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("notes.txt", "MFA is enforced.")
    with pytest.raises(IngestError, match="content must match"):
        parse("x.zip", buf.getvalue())


# Beyond the brief's list: each test below pins a failure that probing the first version turned up.
def _cell(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def _sheet_xlsx(dimension: str | None, rows: dict[int, list[str]]) -> bytes:
    """A real workbook whose first sheet is rewritten by hand: a declared size and rows by row number."""
    wb = openpyxl.Workbook()
    wb.active.append(["placeholder"])
    saved = io.BytesIO()
    wb.save(saved)
    body = "".join(
        f'<row r="{r}">' + "".join(_cell(f"{'ABCDEFG'[i]}{r}", v) for i, v in enumerate(vals)) + "</row>"
        for r, vals in sorted(rows.items())
    )
    dim = f'<dimension ref="{dimension}"/>' if dimension else ""
    sheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"{dim}<sheetData>{body}</sheetData></worksheet>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(saved) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename == "xl/worksheets/sheet1.xml":
                zout.writestr(info.filename, sheet)
            else:
                zout.writestr(info.filename, zin.read(info.filename))
    return out.getvalue()


@pytest.mark.parametrize("dimension", ["A1:B1", "A1", "A1:B3", None])
def test_a_sheet_is_read_in_full_whatever_size_it_declares(dimension: str | None) -> None:
    # openpyxl's read-only mode stops at the declared size; some writers declare A1 for every sheet.
    rows = {1: ["System", "Owner"], **{i: [f"app{i}", "IT"] for i in range(2, 7)}}
    lines = parse("declared.xlsx", _sheet_xlsx(dimension, rows)).lines
    assert [x.text for x in lines] == [f"System: app{i}; Owner: IT" for i in range(2, 7)]


def test_a_sheet_that_claims_a_row_far_below_is_refused_not_walked(monkeypatch: pytest.MonkeyPatch) -> None:
    # Rows are counted as they are read, empty ones included: a tiny file can name row 10**9.
    monkeypatch.setattr(parse_module, "MAX_ROWS", 1_000)
    with pytest.raises(IngestError, match="more than 1,000 rows"):
        parse("far.xlsx", _sheet_xlsx("A1:B5000", {1: ["System", "Owner"], 5000: ["Okta", "IT"]}))


def test_the_row_limit_is_shared_by_all_sheets(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(parse_module, "MAX_ROWS", 1_000)

    def book(sheets: int) -> bytes:
        wb = openpyxl.Workbook()
        for i in range(sheets):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.append(["System", "Owner"])
            ws.cell(row=600, column=1, value="Okta")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    assert parse("one.xlsx", book(1)).lines  # 600 rows are read
    with pytest.raises(IngestError, match="more than 1,000 rows"):
        parse("two.xlsx", book(2))  # 1,200


def test_an_as_of_line_with_an_impossible_date_states_nothing() -> None:
    data = NL.join(["Access log", "As of: 2026-13-45", "System,Last review", "Okta,2026-01-10", ""])
    assert [(x.text, x.as_of) for x in parse("log.csv", data.encode()).lines] == [
        ("Access log", None),
        ("As of: 2026-13-45", None),
        ("System: Okta; Last review: 2026-01-10", date(2026, 1, 10)),
    ]


def test_a_csv_with_old_macintosh_line_endings_is_read() -> None:
    assert [x.text for x in parse("mac.csv", b"System,Status\rOkta,Overdue\r").lines] == [
        "System: Okta; Status: Overdue"
    ]


def test_a_csv_cell_beyond_the_csv_modules_limit_gets_a_readable_reason() -> None:
    with pytest.raises(IngestError, match="20,000 characters"):
        parse("wide.csv", b"System,Note\nOkta," + b"x" * 200_000 + b"\n")


@pytest.mark.parametrize(("name", "member"), [("a.docx", "word/document.xml"), ("a.xlsx", "xl/workbook.xml")])
def test_an_office_file_that_is_not_a_whole_package_is_refused_as_damaged(name: str, member: str) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(member, "<root/>")
    with pytest.raises(IngestError, match="damaged"):
        parse(name, buf.getvalue())


def test_a_docx_with_too_many_paragraphs_is_refused_without_reading_them_all(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(parse_module, "MAX_LINES", 50)
    seen: list[str] = []
    real = parse_module.normalize
    monkeypatch.setattr(parse_module, "normalize", lambda t: seen.append(t) or real(t))
    doc = docx.Document()
    for i in range(500):
        doc.add_paragraph(f"Line {i}.")
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(IngestError, match="50 lines"):
        parse("long.docx", buf.getvalue())
    assert len(seen) <= 52


def test_a_docx_table_with_too_many_rows_is_refused_without_reading_them_all(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(parse_module, "MAX_LINES", 50)
    seen: list[object] = []
    real = parse_module.record_line
    monkeypatch.setattr(parse_module, "record_line", lambda h, v: seen.append(v) or real(h, v))
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "System", "Owner"
    for i in range(500):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = f"app{i}", "IT"
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(IngestError, match="50 lines"):
        parse("table.docx", buf.getvalue())
    assert len(seen) <= 52


def test_a_workbook_with_too_many_lines_across_sheets_stops_reading_early(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(parse_module, "MAX_LINES", 50)
    seen: list[object] = []
    real = parse_module.record_line
    monkeypatch.setattr(parse_module, "record_line", lambda h, v: seen.append(v) or real(h, v))
    wb = openpyxl.Workbook()
    for i in range(5):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.append(["System", "Owner"])
        for n in range(30):  # each sheet is under the limit; two of them are not
            ws.append([f"app{n}", "IT"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(IngestError, match="50 lines"):
        parse("many.xlsx", buf.getvalue())
    assert len(seen) <= 60


# Fix round 1 (review of 1317e38): resource limits, tracked changes, whole-word headers.
def _xlsx_of_empty_far_cells(rows: int) -> bytes:
    wb = openpyxl.Workbook()
    saved = io.BytesIO()
    wb.save(saved)
    sheet = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        + '<row><c r="XFD1"/></row>' * rows
        + "</sheetData></worksheet>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(saved) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = sheet.encode() if info.filename == "xl/worksheets/sheet1.xml" else zin.read(info.filename)
            zout.writestr(info.filename, data)
    return out.getvalue()


def test_the_xlsx_budget_is_charged_per_cell_not_per_row(monkeypatch: pytest.MonkeyPatch) -> None:
    # A row naming column XFD is padded to 16,384 cells: 3 such rows cost 3 row units but 49,152 cells.
    monkeypatch.setattr(parse_module, "MAX_CELLS", 20_000)
    with pytest.raises(IngestError, match="more than 20,000 cells"):
        parse("wide.xlsx", _xlsx_of_empty_far_cells(3))


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_table(rows: list[list[str]], mutate=None) -> bytes:  # type: ignore[no-untyped-def]
    doc = docx.Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, texts in zip(table.rows, rows, strict=True):
        for c, text in zip(r.cells, texts, strict=True):
            c.text = text
    if mutate:
        mutate(table)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_a_docx_cell_spanning_two_columns_fills_both() -> None:
    def span(table) -> None:  # type: ignore[no-untyped-def]
        tr = table.rows[1]._tr
        tr.tc_lst[0].grid_span = 2
        tr.remove(tr.tc_lst[1])

    lines = parse("s.docx", _docx_table([["System", "Owner"], ["Okta", "-"]], span)).lines
    assert [x.text for x in lines] == ["System: Okta; Owner: Okta"]


def test_a_docx_cell_merged_vertically_repeats_the_text_above_it() -> None:
    def merge(table) -> None:  # type: ignore[no-untyped-def]
        table.rows[1]._tr.tc_lst[0].vMerge = "restart"
        table.rows[2]._tr.tc_lst[0].vMerge = "continue"
        table.rows[3]._tr.tc_lst[0].vMerge = "continue"

    data = _docx_table([["Owner", "System"], ["IT", "Okta"], ["", "Jira"], ["", "Slack"]], merge)
    assert [x.text for x in parse("m.docx", data).lines] == [
        "Owner: IT; System: Okta",
        "Owner: IT; System: Jira",
        "Owner: IT; System: Slack",
    ]


def test_a_huge_declared_span_is_bounded_by_the_table_grid() -> None:
    def span(table) -> None:  # type: ignore[no-untyped-def]
        tr = table.rows[1]._tr
        tr.tc_lst[0].grid_span = 100_000
        tr.remove(tr.tc_lst[1])

    lines = parse("h.docx", _docx_table([["System", "Owner"], ["Okta", "-"]], span)).lines
    assert [x.text for x in lines] == ["System: Okta; Owner: Okta"]


def test_a_plain_docx_table_reads_as_before() -> None:
    data = _docx_table([["System", "Owner"], ["Okta", "IT"], ["Jira", ""]])
    assert [x.text for x in parse("p.docx", data).lines] == ["System: Okta; Owner: IT", "System: Jira"]


def test_docx_paragraph_styles_are_resolved_once_per_document(monkeypatch: pytest.MonkeyPatch) -> None:
    from docx.styles.styles import Styles

    calls = {"default": 0, "get_by_id": 0}
    real_default, real_get = Styles.default, Styles.get_by_id

    def default(self, style_type):  # type: ignore[no-untyped-def]
        calls["default"] += 1
        return real_default(self, style_type)

    def get_by_id(self, style_id, style_type):  # type: ignore[no-untyped-def]
        calls["get_by_id"] += 1
        return real_get(self, style_id, style_type)

    doc = docx.Document()
    doc.add_heading("Access", level=1)
    for i in range(40):
        doc.add_paragraph(f"Line {i}.")
        doc.add_paragraph(f"Item {i}.", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    monkeypatch.setattr(Styles, "default", default)  # counted from here: building the file used them too
    monkeypatch.setattr(Styles, "get_by_id", get_by_id)
    lines = parse("styles.docx", buf.getvalue()).lines
    assert [x.text for x in lines if x.kind == "heading"] == ["Access"]
    assert calls["get_by_id"] == 0
    assert calls["default"] <= 1


def test_a_docx_part_larger_than_the_cap_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    buf = io.BytesIO()
    docx.Document().save(buf)
    monkeypatch.setattr(parse_module, "MAX_DOCX_PART", 1_000)
    with pytest.raises(IngestError, match="too large to read"):
        parse("big.docx", buf.getvalue())


def test_a_docx_with_tracked_changes_is_refused() -> None:
    doc = docx.Document()
    p = doc.add_paragraph("MFA is required for admins.")
    p._p.append(parse_xml(f'<w:ins xmlns:w="{_W}" w:id="1" w:author="a"><w:r><w:t> not</w:t></w:r></w:ins>'))
    buf = io.BytesIO()
    doc.save(buf)
    with pytest.raises(IngestError, match="tracked changes. Accept or reject them"):
        parse("t.docx", buf.getvalue())


def test_a_header_that_only_contains_a_date_word_still_dates_the_row() -> None:
    data = "System,Overdue since" + NL + "Okta,2026-01-10" + NL
    assert parse("o.csv", data.encode()).lines[0].as_of == date(2026, 1, 10)

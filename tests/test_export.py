import csv
import io
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import Engine

from app.api.schemas import Mapping
from app.export import DRAFT, ExportRow, export_csv, export_xlsx
from tests.apiclient import visitor

SAMPLES = Path(__file__).resolve().parent.parent / "data" / "questionnaires"
VSQ = Mapping(
    sheet="Questionnaire",
    header_row=5,
    id_col="A",
    question_col="C",
    answer_col="D",
    comments_col="E",
    topic_col="B",
)
ROWS = [
    ExportRow(
        7,
        "verified",
        "Yes",
        "Yes. The policy says so.",
        True,
        ["information-security-policy.docx line 4"],
        [],
    ),
    ExportRow(
        8, "partial", "Partial", "Partly.", False, ["policy.docx line 2"], ["Scope: internal systems only."]
    ),
    ExportRow(9, "unknown", None, "", False, [], []),
]
PAYLOADS = [
    '=HYPERLINK("https://evil.example","Yes")',
    "@SUM(1+1)",
    "+cmd|' /C calc'!A0",
    "-1+1",
    "\t=1+1",
    "\r=1+1",
    "\n=1+1",
    "=1+1",
]


def test_the_xlsx_keeps_its_formatting_and_gets_three_columns() -> None:
    original = (SAMPLES / "vsq-a.xlsx").read_bytes()
    before = openpyxl.load_workbook(io.BytesIO(original))["Questionnaire"]
    wb = openpyxl.load_workbook(io.BytesIO(export_xlsx(original, VSQ, ROWS)))
    ws = wb["Questionnaire"]
    assert [ws.cell(5, c).value for c in (6, 7, 8)] == ["Status", "Sources", "Notes"]
    assert (ws["D7"].value, ws["E7"].value, ws["F7"].value) == (
        "Yes",
        "Yes. The policy says so.",
        "Verified · Approved",
    )
    assert ws["G7"].value == "information-security-policy.docx line 4"
    assert ws["F8"].value == f"Partial · {DRAFT}" and ws["F9"].value == f"Unknown · {DRAFT}"
    # Yes/No/N/A validation without "Partial": the value moves to Notes, the answer cell stays empty.
    assert ws["D8"].value is None and "Answer: Partial" in ws["H8"].value
    # Formatting kept, cell by cell where it matters.
    assert {str(r) for r in ws.merged_cells.ranges} == {str(r) for r in before.merged_cells.ranges}
    assert len(ws.data_validations.dataValidation) == len(before.data_validations.dataValidation)
    assert ws.column_dimensions["C"].width == before.column_dimensions["C"].width
    assert ws["C7"].font.b == before["C7"].font.b and ws["A5"].font.b == before["A5"].font.b


def test_the_csv_keeps_its_columns_and_delimiter() -> None:
    original = b"Question;Answer;Comment\r\nIs MFA on?;;\r\nAre logs kept?;;\r\n"
    mapping = Mapping(
        sheet=None, header_row=1, id_col=None, question_col="A", answer_col="B", comments_col="C"
    )
    out = export_csv(original, mapping, [ExportRow(2, "verified", "Yes", "Yes.", False, ["p.md line 1"], [])])
    rows = list(csv.reader(io.StringIO(out.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["Question", "Answer", "Comment", "Status", "Sources", "Notes"]
    assert rows[1] == ["Is MFA on?", "Yes", "Yes.", f"Verified · {DRAFT}", "p.md line 1", ""]
    assert rows[2][:3] == ["Are logs kept?", "", ""]


@pytest.mark.parametrize("payload", PAYLOADS)
def test_the_xlsx_cells_are_inert_text(payload: str) -> None:
    # Adversary 1 I1: openpyxl turns a string starting with = into a formula; every cell must stay a string.
    original = (SAMPLES / "vsq-a.xlsx").read_bytes()
    row = ExportRow(7, "verified", "Yes", payload, False, [payload], [payload])
    ws = openpyxl.load_workbook(io.BytesIO(export_xlsx(original, VSQ, [row])))["Questionnaire"]
    for ref in ("E7", "G7", "H7"):
        assert (ws[ref].data_type, ws[ref].value) == ("s", payload), ref
    # A value that is itself the payload (an answer column without a validation list).
    mapping = Mapping(
        sheet="Questionnaire", header_row=5, id_col="A", question_col="C", answer_col="E", comments_col="D"
    )
    row = ExportRow(7, "unknown", None, payload, False, [], [])
    ws = openpyxl.load_workbook(io.BytesIO(export_xlsx(original, mapping, [row])))["Questionnaire"]
    assert ws["E7"].data_type == "s" and ws["E7"].value == payload


@pytest.mark.parametrize("payload", PAYLOADS)
def test_the_csv_cells_are_inert_text(payload: str) -> None:
    original = b"Question,Answer,Comment\r\nIs MFA on?,,\r\n"
    mapping = Mapping(
        sheet=None, header_row=1, id_col=None, question_col="A", answer_col="B", comments_col="C"
    )
    out = export_csv(original, mapping, [ExportRow(2, "unknown", None, payload, False, [payload], [payload])])
    rows = list(csv.reader(io.StringIO(out.decode("utf-8-sig"), newline="")))
    for cell in rows[1][1:]:
        assert cell.startswith("'") or cell[:1] not in tuple("=+-@\t\r\n"), cell
    assert rows[1][1] == "'" + payload  # the answer cell
    assert rows[1][4] == "'" + payload  # sources


def test_the_endpoint_refuses_an_id_that_is_not_this_workspaces_run(db: Engine) -> None:
    client, _ = visitor(db)
    q = client.post("/api/questionnaires/sample/mvsp-b").json()
    assert client.get(f"/api/runs/{q['id']}/export").status_code == 404  # a questionnaire id is not a run


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("MVSP.csv", 'attachment; filename="MVSP-filled.csv"'),
        ('a"b\r\nSet-Cookie: x=1.xlsx', 'attachment; filename="a_b_Set-Cookie_x_1-filled.xlsx"'),
        ("büro.xlsx", 'attachment; filename="b_ro-filled.xlsx"'),
        ("\r\n.xlsx", 'attachment; filename="questionnaire-filled.xlsx"'),
    ],
)
def test_the_download_header_is_safe(filename: str, expected: str) -> None:
    from urllib.parse import unquote

    from app.api.export import disposition

    head = disposition(filename, filename.rsplit(".", 1)[1])
    assert head.startswith(expected)
    assert not any(c in head for c in "\r\n") and head.isascii()
    if head != expected:
        assert head.count('"') == 2 and unquote(head.split("UTF-8''")[1]).startswith(
            filename.rsplit(".", 1)[0]
        )

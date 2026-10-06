import io
import json
import threading
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import Engine, func, select, update
from sqlalchemy.orm import Session

from app import csf, runs
from app.api.deps import get_llm
from app.api.runs import summary
from app.db.models import Answer, Document, DocumentLine, Item, Questionnaire, Run, RunItem, Workspace
from app.export import FOOTER
from app.main import app
from app.services.ip_limits import LIMITS
from tests import factories as f
from tests.apiclient import visitor
from tests.fakes import ByStepLLM

LINE = "Backups of data are created, protected, maintained and tested every day."
YES = json.dumps({"passages": [{"passage": 1, "stance": "yes", "quote": LINE, "note": "states it"}]})
SAID = "Our cybersecurity policy is established and communicated to all staff."
FILL = json.dumps({"passages": [{"passage": 1, "stance": "yes", "quote": SAID, "note": "says it"}]})
NO_FILL = json.dumps({"passages": []})


def _finish(client: TestClient, run_id: str, llm: object) -> None:
    app.dependency_overrides[get_llm] = lambda: llm
    try:
        for _ in range(60):
            if client.get(f"/api/runs/{run_id}").json()["status"] != "running":
                return
            assert client.post(f"/api/runs/{run_id}/step").status_code == 200
        raise AssertionError("the run did not finish")
    finally:
        app.dependency_overrides.pop(get_llm, None)


def _policy(s: Session, ws_id: uuid.UUID, status: str = "final") -> None:
    """One stored policy line, as an upload leaves it (document, line, chunk)."""
    doc = f.document(s, s.get_one(Workspace, ws_id), filename="backup-policy.docx", status=status)
    s.add(DocumentLine(document_id=doc.id, n=1, text=LINE))
    f.chunk(s, doc, line_start=1, line_end=1, text=LINE)
    s.commit()


def test_the_view_lists_every_outcome_and_labels_only_what_was_checked(db: Engine) -> None:
    client, _ = visitor(db)
    before = client.get("/api/gap/core").json()
    assert before["run"] is None and len(before["rows"]) == 106
    assert {r["label"] for r in before["rows"]} == {None}
    tiers = [r["tier"] for r in before["rows"]]
    assert (tiers.count("checked"), tiers.count("ask"), tiers.count("not_checked")) == (31, 5, 70)
    assert before["rows"][0]["outcome"] == csf.framework().outcomes[0].outcome  # NIST's text, verbatim
    with Session(db) as s:
        assert s.scalar(select(func.count()).select_from(Questionnaire)) == 0  # the view writes nothing
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({}))  # no documents: no part has a passage, nothing is called
    after = client.get("/api/gap/core").json()
    rows = {r["csf_id"]: r for r in after["rows"]}
    assert (rows["PR.DS-11"]["label"], rows["PR.DS-11"]["explanation"]) == (
        "gap",
        "No evidence: parts 1, 2, 3, 4.",
    )
    assert (rows["GV.RM-02"]["label"], rows["GV.RM-02"]["explanation"]) == ("not_answered", None)
    assert all(
        r["label"] is None and r["item_id"] is None for r in after["rows"] if r["tier"] == "not_checked"
    )
    assert after["run"]["status"] == "done" and after["controls_url"] == csf.CONTROLS_URL


def test_a_function_scope_lists_its_own_outcomes_and_runs_only_them(db: Engine) -> None:
    client, _ = visitor(db)
    rows = client.get("/api/gap/recover").json()["rows"]
    assert {r["function"] for r in rows} == {"Recover"}
    assert client.post("/api/gap/recover/run").json()["total"] == 1  # RC.RP-01: Recover's one checked outcome
    assert client.get("/api/gap/everything").status_code == 422


def test_starting_twice_continues_the_same_run_and_the_questionnaire_stays_built_in(db: Engine) -> None:
    client, _ = visitor(db)
    first = client.post("/api/gap/core/run").json()
    again = client.post("/api/gap/core/run").json()
    assert (again["id"], again["status"]) == (first["id"], "running")
    assert client.get("/api/questionnaires").json() == []  # Plan 3 Ruling 5: never listed or counted


def test_starting_counts_under_the_run_limit(db: Engine, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(LIMITS, "run", (1, timedelta(hours=1)))
    client, _ = visitor(db)
    assert client.post("/api/gap/core/run").status_code == 200
    res = client.post("/api/gap/govern/run")
    assert res.status_code == 429 and res.headers["Retry-After"]


def test_another_workspace_sees_none_of_it(db: Engine) -> None:
    client, _ = visitor(db)
    client.post("/api/gap/core/run")
    other, _ = visitor(db)
    assert other.get("/api/gap/core").json()["run"] is None


def test_an_outcomes_evidence_lists_its_parts_with_their_own_citations(db: Engine) -> None:
    client, ws_id = visitor(db)
    with Session(db) as s:
        _policy(s, ws_id, status="draft")
    run = client.post("/api/gap/protect/run").json()
    _finish(client, run["id"], ByStepLLM({"stance": YES}))
    row = next(r for r in client.get("/api/gap/protect").json()["rows"] if r["csf_id"] == "PR.DS-11")
    assert row["label"] == "partly_covered"  # every part rests on a draft
    detail = client.get(f"/api/answers/{row['answer_id']}").json()
    parts = detail["parts"]
    assert [p["n"] for p in parts] == [1, 2, 3, 4]
    assert [p["question"] for p in parts] == list(csf.framework().get("PR.DS-11").parts)
    assert {p["label"] for p in parts} == {"partly_covered"} and not any(p["from_statement"] for p in parts)
    cite = parts[0]["citations"][0]
    assert (cite["quote"], cite["status"], cite["found_in_source"]) == (LINE, "draft", True)  # carry d


def test_the_gap_report_downloads_for_a_gap_run(db: Engine) -> None:
    client, _ = visitor(db)
    run = client.post("/api/gap/recover/run").json()
    _finish(client, run["id"], ByStepLLM({}))
    res = client.get(f"/api/runs/{run['id']}/export")
    assert res.status_code == 200
    assert res.headers["content-disposition"] == 'attachment; filename="csf-2.0-recover-gap-report.xlsx"'
    ws = openpyxl.load_workbook(io.BytesIO(res.content))["Gap report"]
    recover = [o for o in csf.framework().outcomes if o.function == "Recover"]
    assert [ws.cell(n, 1).value for n in range(3, 3 + len(recover))] == [o.id for o in recover]
    assert ws.cell(len(recover) + 4, 1).value == FOOTER
    assert (ws["B1"].value, ws["C1"].value) == ("Scope: recover", f"Run date: {run['started_at'][:10]}")


def test_a_questionnaire_export_carries_the_latest_gap_sheet_only_when_one_exists(db: Engine) -> None:
    client, _ = visitor(db)
    q = client.post("/api/questionnaires/sample/vsq-a").json()
    qrun = client.post(f"/api/questionnaires/{q['id']}/runs").json()  # no step needed: an empty run exports
    before = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/runs/{qrun['id']}/export").content))
    assert "Gap report" not in before.sheetnames  # no gap check yet: the export is unchanged
    gap = client.post("/api/gap/recover/run").json()
    _finish(client, gap["id"], ByStepLLM({}))
    after = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/runs/{qrun['id']}/export").content))
    assert after.sheetnames == [*before.sheetnames, "Gap report"]
    ws = after["Gap report"]
    assert (ws["A1"].value, ws["B1"].value) == ("Possible gap — review it", "Scope: recover")
    assert ws["A3"].value == "RC.RP-01" and ws["D3"].value == "Gap"


def test_two_first_presses_at_once_create_one_run(db: Engine) -> None:
    client, ws_id = visitor(db)
    other = TestClient(app, cookies=client.cookies)
    results: list[dict] = []  # type: ignore[type-arg]
    gate = threading.Barrier(2)

    def press(c: TestClient) -> None:
        gate.wait()
        results.append(c.post("/api/gap/recover/run").json())

    threads = [threading.Thread(target=press, args=(c,)) for c in (client, other)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    with Session(db) as s:
        assert s.scalar(select(func.count()).select_from(Run).where(Run.workspace_id == ws_id)) == 1
    assert len(results) == 2 and results[0]["id"] == results[1]["id"]


def _answer(db: Engine, csf_id: str) -> uuid.UUID:
    with Session(db) as s:
        return s.scalars(
            select(Answer.id).join(Item, Item.id == Answer.item_id).where(Item.csf_id == csf_id)
        ).one()


def test_not_applicable_and_failed_outcomes_carry_no_gap_label(db: Engine) -> None:
    client, _ = visitor(db)
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({}))
    na_checked, na_ask, failed = _answer(db, "PR.DS-11"), _answer(db, "GV.RM-02"), _answer(db, "PR.AA-01")
    for aid in (na_checked, na_ask):
        assert (
            client.post(f"/api/answers/{aid}/not-applicable", json={"reason": "not ours"}).status_code == 200
        )
    with Session(db) as s:
        s.execute(update(Answer).where(Answer.id == failed).values(label="unknown", text=runs.FAILED_TEXT))
        s.commit()
    rows = {r["csf_id"]: r for r in client.get("/api/gap/core").json()["rows"]}
    for cid in ("PR.DS-11", "GV.RM-02"):
        assert (rows[cid]["label"], rows[cid]["not_applicable"]) == (None, True)
        assert rows[cid]["explanation"] == "Not applicable: not ours"
    assert (rows["PR.AA-01"]["label"], rows["PR.AA-01"]["not_applicable"]) == (None, False)
    assert (
        rows["PR.AA-01"]["explanation"] == "Not checked: the model call failed twice. Press r to check again."
    )
    assert client.get(f"/api/answers/{na_checked}").json()["parts"] == []  # no stale parts for N/A (M2)
    sheet = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/runs/{run['id']}/export").content))[
        "Gap report"
    ]
    words = {sheet.cell(n, 1).value: sheet.cell(n, 4).value for n in range(3, 109)}
    assert (words["PR.DS-11"], words["GV.RM-02"], words["PR.AA-01"]) == (
        "Not applicable",
        "Not applicable",
        "Failed",
    )
    failed_row = next(n for n in range(3, 109) if sheet.cell(n, 1).value == "PR.AA-01")
    assert sheet.cell(failed_row, 5).value == "Not checked: the model call failed twice."  # no UI hint


def test_a_failed_outcome_shows_no_parts_and_an_incomplete_one_neither(db: Engine) -> None:
    client, _ = visitor(db)
    run = client.post("/api/gap/protect/run").json()
    _finish(client, run["id"], ByStepLLM({}))
    aid = _answer(db, "PR.DS-11")
    assert client.get(f"/api/answers/{aid}").json()["parts"] != []
    with Session(db) as s:  # a part missing from the stored parts (adversary-1 M2)
        key = (RunItem.run_id == uuid.UUID(run["id"]), RunItem.item_id == s.get_one(Answer, aid).item_id)
        parts = s.scalars(select(RunItem.parts).where(*key)).one()
        parts.pop(next(iter(parts)))
        s.execute(update(RunItem).where(*key).values(parts=parts))
        s.commit()
    assert client.get(f"/api/answers/{aid}").json()["parts"] == []
    with Session(db) as s:  # a failed outcome keeps parts, but its answer is not a result of them
        s.execute(update(Answer).where(Answer.id == aid).values(label="unknown", text=runs.FAILED_TEXT))
        s.commit()
    assert client.get(f"/api/answers/{aid}").json()["parts"] == []


def test_bulk_approve_skips_every_gap_check_outcome(db: Engine) -> None:
    client, ws_id = visitor(db)
    run = client.post("/api/gap/recover/run").json()
    _finish(client, run["id"], ByStepLLM({}))
    aid = _answer(db, "RC.RP-01")
    with Session(db) as s:  # N3: a Not met outcome (verified, No) is not approved in bulk
        s.execute(
            update(Answer)
            .where(Answer.id == aid)
            .values(label="verified", value="No", text="x", citations=[f.CITATION])
        )
        s.commit()
    assert client.post(f"/api/runs/{run['id']}/approve-verified").json()["approved"] == 0
    with Session(db) as s:
        s.execute(update(Answer).where(Answer.id == aid).values(value="Yes"))
        s.commit()
    # adversary-2 M5: a Covered outcome is not approved either; an approval would keep Check again off it
    assert client.post(f"/api/runs/{run['id']}/approve-verified").json()["approved"] == 0
    q = client.post(
        "/api/questionnaires/sample/vsq-a"
    ).json()  # a questionnaire's verified No is still approved
    qrun = client.post(f"/api/questionnaires/{q['id']}/runs").json()
    with Session(db) as s:
        item = s.scalars(select(Item).where(Item.questionnaire_id == uuid.UUID(q["id"]))).first()
        assert item is not None
        f.answer(
            s,
            s.get_one(Run, uuid.UUID(qrun["id"])),
            item,
            label="verified",
            value="No",
            text="x",
            citations=[f.CITATION],
        )
        s.commit()
    assert client.post(f"/api/runs/{qrun['id']}/approve-verified").json()["approved"] == 1


def _gap_run_id(client: TestClient, scope: str) -> str:
    return str(client.post(f"/api/gap/{scope}/run").json()["id"])


def test_the_questionnaire_export_carries_the_latest_finished_gap_check(db: Engine) -> None:
    client, _ = visitor(db)
    q = client.post("/api/questionnaires/sample/vsq-a").json()
    qrun = client.post(f"/api/questionnaires/{q['id']}/runs").json()

    def sheet() -> Any:
        body = client.get(f"/api/runs/{qrun['id']}/export").content
        return openpyxl.load_workbook(io.BytesIO(body))["Gap report"]

    recover = _gap_run_id(client, "recover")
    _finish(client, recover, ByStepLLM({}))
    govern = _gap_run_id(client, "govern")  # left running: skipped
    assert sheet()["B1"].value == "Scope: recover"
    _finish(client, govern, ByStepLLM({}))
    assert sheet()["B1"].value == "Scope: govern"  # the latest finished wins over the older one
    with Session(db) as s:
        s.execute(
            update(Run)
            .where(Run.id == uuid.UUID(recover))
            .values(finished_at=datetime(2030, 1, 2, tzinfo=UTC))
        )
        s.commit()
    ws = sheet()
    assert (ws["B1"].value, ws["C1"].value) == ("Scope: recover", "Run date: 2030-01-02")


def test_check_again_after_an_upload_reopens_the_outcomes_the_new_document_reaches(db: Engine) -> None:
    client, ws_id = visitor(db)
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({}))  # no documents: every Checked outcome is Gap
    with Session(db) as s:
        _policy(s, ws_id)  # an upload
        # adversary-1 I4 (d): exactly the parts the new line reaches are judged, nothing else
        reached = [
            p.key
            for o in csf.in_scope("core")
            for p in csf.part_inputs(o)
            if csf.evidence(s, ws_id, p).passages
        ]
    llm = ByStepLLM({"stance": YES})
    again = client.post("/api/gap/core/run").json()
    assert (again["id"], again["status"]) == (run["id"], "running")
    _finish(client, run["id"], llm)
    rows = {r["csf_id"]: r for r in client.get("/api/gap/core").json()["rows"]}
    assert rows["PR.DS-11"]["label"] == "covered"
    assert rows["GV.RM-02"]["label"] == "not_answered"  # Ask me is never re-checked
    assert "PR.DS-11#1" in reached and len(reached) < 73
    assert sorted(r.item_id for r in llm.requests) == sorted(reached)
    calls = len(llm.requests)
    assert client.post("/api/gap/core/run").json()["status"] == "done"  # nothing changed since
    assert len(llm.requests) == calls


def test_on_a_one_document_workspace_an_ask_me_answer_alone_reopens_nothing(db: Engine) -> None:
    """One policy line only: the sample pack's case is the test below (adversary-2 I2)."""
    client, ws_id = visitor(db)
    with Session(db) as s:
        _policy(s, ws_id)
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({"stance": YES}))
    question = next(
        q for q in client.get(f"/api/runs/{run['id']}/questions").json() if q["codes"] == ["GV.RM-02"]
    )
    llm = ByStepLLM({"recheck": json.dumps({"passages": []})})  # the fills found nothing
    app.dependency_overrides[get_llm] = lambda: llm
    try:
        text = "Backups of data are created, protected, maintained and tested; our risk appetite is low."
        assert client.post(f"/api/questions/{question['id']}/answer", json={"text": text}).status_code == 200
    finally:
        app.dependency_overrides.pop(get_llm, None)
    llm.requests.clear()
    assert client.post("/api/gap/core/run").json()["status"] == "done"
    _finish(client, run["id"], llm)
    assert llm.requests == []


def _ask(client: TestClient, run_id: str, csf_id: str, text: str, reply: str) -> Any:
    """Answer an Ask-me outcome's question, the re-check replying `reply`."""
    question = next(q for q in client.get(f"/api/runs/{run_id}/questions").json() if q["codes"] == [csf_id])
    app.dependency_overrides[get_llm] = lambda: ByStepLLM({"recheck": reply})
    try:
        r = client.post(f"/api/questions/{question['id']}/answer", json={"text": text})
    finally:
        app.dependency_overrides.pop(get_llm, None)
    assert r.status_code == 200
    return r.json()


def test_the_visitors_answer_never_passes_as_document_evidence(db: Engine) -> None:
    """adversary-2 I1, the probe replayed: a statement's kind cannot be patched, and a statement whose kind
    changed anyway is still left out of every Checked part's evidence and still marked "(your answer)"."""
    client, ws_id = visitor(db)
    with Session(db) as s:
        _policy(s, ws_id)
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({"stance": YES}))
    found = _ask(client, run["id"], "GV.RM-02", SAID, FILL)["suggestions"]
    fill = next(x for x in found if x["code"] == "GV.PO-01" and x["part"] == 1)
    assert client.post(f"/api/suggestions/{fill['id']}/accept").status_code == 200
    stmt = next(d for d in client.get("/api/documents").json() if d["source"] == "statement")
    assert client.patch(f"/api/documents/{stmt['id']}", json={"kind": "policy"}).status_code == 409
    with Session(db) as s:  # as if laundered before the guard: nothing keys on a statement's kind
        s.execute(update(Document).where(Document.id == uuid.UUID(stmt["id"])).values(kind="policy"))
        s.commit()
        o = csf.framework().get("GV.PO-01")
        for part in csf.part_inputs(o):
            assert stmt["filename"] not in {p.doc.filename for p in csf.evidence(s, ws_id, part).passages}
    assert client.post("/api/gap/core/run").json()["status"] == "done"  # nothing re-opened, nothing paid
    row = next(r for r in client.get("/api/gap/core").json()["rows"] if r["csf_id"] == "GV.PO-01")
    assert row["label"] not in ("covered", "confirmed_by_you")  # one part from the answer, the rest Gaps
    sheet = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/runs/{run['id']}/export").content))[
        "Gap report"
    ]
    quotes = next(
        sheet.cell(n, 6).value for n in range(3, sheet.max_row + 1) if sheet.cell(n, 1).value == "GV.PO-01"
    )
    assert f'"{SAID}" ({stmt["filename"]} line 1) (your answer)' in quotes


def test_on_the_sample_pack_an_ask_me_answer_alone_reopens_nothing(db: Engine) -> None:
    """adversary-2 I2: on the real 23 documents, storing the E2E's own answer moves no part's evidence (the
    IDF and the chunk total leave statements out too), so Check again stays done."""
    client, _ = visitor(db)
    assert client.post("/api/documents/sample").status_code == 201
    run = client.post("/api/gap/core/run").json()
    _finish(client, run["id"], ByStepLLM({"stance": YES}))
    answer = (
        "Yes. The board approved a cybersecurity risk appetite statement, and the security team shares it "
        "with every new hire."
    )  # web/e2e/gap.spec.ts
    _ask(client, run["id"], "GV.RM-02", answer, NO_FILL)
    assert client.post("/api/gap/core/run").json()["status"] == "done"
    row = next(r for r in client.get("/api/gap/core").json()["rows"] if r["csf_id"] == "GV.RM-02")
    # adversary-2 M8: the visitor's own words are said to be theirs
    assert (row["label"], row["explanation"]) == ("confirmed_by_you", f"Your answer: {answer}")


def test_a_gap_outcome_is_never_edited_or_approved_from_the_run_view(db: Engine) -> None:
    """adversary-2 M5 and M6: an edit or approval would keep Check again off the outcome and put text code
    did not write into the sheet, so both are refused for a gap check's outcome."""
    client, _ = visitor(db)
    run = client.post("/api/gap/recover/run").json()
    _finish(client, run["id"], ByStepLLM({}))
    aid = _answer(db, "RC.RP-01")
    assert client.patch(f"/api/answers/{aid}", json={"text": "We test it."}).status_code == 409
    approve = client.post(f"/api/answers/{aid}/approve")
    assert approve.status_code == 409 and "question" not in approve.json()["detail"]
    with Session(db) as s:
        a = s.get_one(Answer, aid)
        assert (a.edited, a.approved_at, a.label) == (False, None, "unknown")


def test_a_running_gap_check_is_not_exported(db: Engine) -> None:
    """adversary-2 M3: re-opened rows would read "Not run yet" under the first run's date."""
    client, _ = visitor(db)
    run = client.post("/api/gap/recover/run").json()
    assert client.get(f"/api/runs/{run['id']}/export").status_code == 409
    _finish(client, run["id"], ByStepLLM({}))
    assert client.get(f"/api/runs/{run['id']}/export").status_code == 200


def test_an_outcome_confirmed_through_a_fill_counts_every_cited_document() -> None:
    """adversary-2 M2: Confirmed by you through a part fill cites documents and the statement."""

    def confirmed(*docs: str) -> Answer:
        cites = [{"document_id": d} for d in docs]
        return Answer(
            id=uuid.uuid4(),
            item_id=uuid.uuid4(),
            label="user_confirmed",
            text="x",
            confidence=1.0,
            edited=False,
            citations=cites,
        )

    assert summary(confirmed("doc-1", "doc-2", "answer-1")).sources == 3
    assert summary(confirmed()).sources == 1  # an Ask-me confirmation: the statement

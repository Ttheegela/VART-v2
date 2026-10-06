"""Adversary checkpoint 3 fixes for the inputs lane (C1, I1-I4, M1-M3, M5, M6)."""

import csv
import io
import zipfile

import openpyxl
import pytest
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from app.api.schemas import Mapping
from app.db.models import Workspace
from app.export import ExportRow, export_csv, export_xlsx
from app.ingest.parse import MAX_TEXT_CHARS, IngestError, parse
from app.questionnaires import MAX_QUESTION, MAX_TOPIC, read_items, read_sheets
from tests import factories as f
from tests.apiclient import visitor

NO_ANSWER = Mapping(sheet=None, header_row=1, id_col="A", question_col="B", answer_col="C", comments_col=None)


def _xlsx(rows: list[list[str]], title: str = "S") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_c1_a_csv_without_an_answer_column_exports() -> None:
    out = export_csv(
        b"ID,Question\r\n1,Is MFA on?\r\n", NO_ANSWER, [ExportRow(2, "verified", "Yes", "Yes.", True, [], [])]
    )
    rows = list(csv.reader(io.StringIO(out.decode())))
    assert rows[0] == ["ID", "Question", "", "Status", "Sources", "Notes"]
    assert rows[1][:4] == ["1", "Is MFA on?", "Yes", "Verified · Approved"]


def test_c1_a_comments_column_past_the_csv_width_exports() -> None:
    m = Mapping(sheet=None, header_row=1, id_col=None, question_col="A", answer_col="B", comments_col="Z")
    out = export_csv(b"Q,A\r\nx,\r\n", m, [ExportRow(2, "unknown", None, "t", False, [], [])])
    assert len(list(csv.reader(io.StringIO(out.decode())))[1]) == 26 + 3


def test_c1_an_xlsx_without_an_answer_column_keeps_the_answer_and_the_status_apart() -> None:
    original = _xlsx([["ID", "Question"], ["1", "Is MFA on?"]])
    ws = openpyxl.load_workbook(
        io.BytesIO(export_xlsx(original, NO_ANSWER, [ExportRow(2, "verified", "Yes", "Yes.", True, [], [])]))
    ).active
    assert [c.value for c in ws[2]][:4] == ["1", "Is MFA on?", "Yes", "Verified · Approved"]
    assert ws["D1"].value == "Status"


def test_i1_a_workbook_that_is_large_once_unpacked_is_refused() -> None:
    sheet = "<worksheet><sheetData>" + "<row><c><v>1</v></c></row>" * 400_000 + "</sheetData></worksheet>"
    buf = io.BytesIO(_xlsx([["a"]]))
    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in zin.infolist():
            zout.writestr(
                i.filename, sheet if i.filename == "xl/worksheets/sheet1.xml" else zin.read(i.filename)
            )
    assert len(out.getvalue()) < 1_000_000
    with pytest.raises(IngestError, match="too large"):
        read_sheets("big.xlsx", out.getvalue())


def test_i2_a_long_question_or_topic_is_refused() -> None:
    m = Mapping(
        sheet=None,
        header_row=1,
        id_col=None,
        question_col="A",
        answer_col="B",
        comments_col=None,
        topic_col="C",
    )
    _, sheets = read_sheets("q.csv", f"Question,Answer,Topic\r\n{'x' * (MAX_QUESTION + 1)},,t\r\n".encode())
    with pytest.raises(IngestError, match="question is longer"):
        read_items(sheets, m)
    _, sheets = read_sheets("q.csv", f"Question,Answer,Topic\r\nq,,{'t' * (MAX_TOPIC + 1)}\r\n".encode())
    with pytest.raises(IngestError, match="topic is longer"):
        read_items(sheets, m)


def test_i4_a_document_over_the_text_cap_is_refused() -> None:
    line = "a" * 10_000 + "\n\n"
    with pytest.raises(IngestError, match="characters of text"):
        parse("big.txt", (line * (MAX_TEXT_CHARS // 10_000 + 2)).encode())


def test_m3_a_hidden_sheet_is_skipped() -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Shown"
    wb.create_sheet("Hidden").sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    _, sheets = read_sheets("q.xlsx", out.getvalue())
    assert [s.name for s in sheets] == ["Shown"]


def test_m6_a_csv_over_the_export_caps_is_refused_at_import() -> None:
    with pytest.raises(IngestError):
        read_sheets("q.csv", ("Q,A\r\n" + "," * 60 + "\r\n").encode())


def test_m5_a_validation_range_is_not_an_inline_list() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Q", "A"])
    ws.append(["x", None])
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(type="list", formula1="=$F$1:$F$3")
    dv.add("B2")
    ws.add_data_validation(dv)
    out = io.BytesIO()
    wb.save(out)
    m = Mapping(sheet=None, header_row=1, id_col=None, question_col="A", answer_col="B", comments_col=None)
    got = openpyxl.load_workbook(
        io.BytesIO(export_xlsx(out.getvalue(), m, [ExportRow(2, "verified", "Yes", "Yes.", True, [], [])]))
    ).active
    assert got["B2"].value == "Yes"


def test_i3_a_document_cannot_be_deleted_while_a_run_is_running(db: Engine) -> None:
    from tests.test_api_documents import _upload

    client, ws_id = visitor(db)
    doc = _upload(client).json()
    with Session(db) as s:
        q = f.questionnaire(s, s.get_one(Workspace, ws_id))
        run = f.run(s, q)
        run.status = "running"
        s.commit()
    r = client.delete(f"/api/documents/{doc['id']}")
    assert (r.status_code, r.json()) == (409, {"detail": "A run is in progress; wait for it to finish."})


def test_m1_a_huge_lines_offset_is_a_422(db: Engine) -> None:
    from tests.test_api_documents import _upload

    client, _ = visitor(db)
    doc = _upload(client).json()
    assert client.get(f"/api/documents/{doc['id']}/lines?from=99999999999").status_code == 422


def test_m2_a_questionnaire_file_name_is_normalised(db: Engine) -> None:
    client, _ = visitor(db)
    name = "q\x00‮gnp‬" + "a" * 300 + ".csv"
    r = client.post(
        "/api/questionnaires",
        files={"file": (name, io.BytesIO(b"Question,Answer\r\nIs MFA on?,\r\n"), "text/csv")},
    )
    assert r.status_code == 201, r.text
    got = r.json()["filename"]
    assert len(got) <= 255 and got.endswith(".csv") and "\x00" not in got and "‮" not in got

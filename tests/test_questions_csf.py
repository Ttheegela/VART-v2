import io
import json
from collections.abc import Iterator
from datetime import date

import openpyxl
import pytest
from sqlalchemy import Engine, delete, func, select, update
from sqlalchemy.orm import Session

from app import csf, export, runs
from app import questions as qs
from app.api.gap import gap_sheet
from app.api.questions import _suggestion_out
from app.db.models import Answer, Item, Questionnaire, RunItem, SuggestedFill
from tests import factories as f
from tests.fakes import ByStepLLM

MODELS = {"stance": "m/s", "draft": "m/d", "classify": "m/c", "recheck": "m/r", "judge": "m/j"}
TODAY = date(2026, 10, 6)
SAID = "Our cybersecurity policy is established and communicated to all staff."
REPLY = json.dumps({"passages": [{"passage": 1, "stance": "yes", "quote": SAID, "note": "x"}]})


@pytest.fixture
def s(db: Engine) -> Iterator[Session]:
    with Session(db) as session:
        yield session


def _core_done(s: Session):  # type: ignore[no-untyped-def]
    """A core gap run over no documents, done: every Checked outcome is Gap, every Ask-me one Not answered."""
    ws = f.workspace(s)
    s.commit()
    q = csf.questionnaire_for(s, ws.id, "core")
    run = runs.create_run(s, ws.id, q.id, MODELS)
    while runs.step(s, ws.id, run.id, ByStepLLM({}), MODELS):
        pass
    return ws, q, run


def _code(s: Session, item_ids: list) -> str:  # type: ignore[type-arg]
    return s.get_one(Item, item_ids[0]).csf_id or ""


def _item(s: Session, q_id: object, csf_id: str) -> Item:
    return s.scalars(select(Item).where(Item.questionnaire_id == q_id, Item.csf_id == csf_id)).one()


def _answer_id(s: Session, run_id: object, item_id: object) -> object:
    return s.scalar(select(Answer.id).where(Answer.run_id == run_id, Answer.item_id == item_id))


def _govern_fills(s: Session):  # type: ignore[no-untyped-def]
    ws, q, run = _core_done(s)
    question = next(x for x in qs.ensure_questions(s, ws.id, run.id) if _code(s, x.item_ids) == "GV.RM-02")
    llm = ByStepLLM({"recheck": REPLY})
    _, _, found = qs.answer_question(s, ws.id, question.id, SAID, llm, MODELS, TODAY)
    return ws, q, run, llm, found


def test_a_gap_run_asks_its_ask_me_outcomes_only(s: Session) -> None:
    ws, q, run = _core_done(s)
    asked = sorted(_code(s, x.item_ids) for x in qs.ensure_questions(s, ws.id, run.id))
    assert asked == [
        "GV.OC-03",
        "GV.OV-01",
        "GV.RM-02",
        "GV.RR-02",
        "GV.SC-01",
    ]  # no Checked Gap is a question


def test_an_ask_me_answer_is_redacted_stored_and_confirmed(s: Session) -> None:
    ws, q, run = _core_done(s)
    question = next(x for x in qs.ensure_questions(s, ws.id, run.id) if _code(s, x.item_ids) == "GV.RM-02")
    text = "Dana Ortiz signed our risk appetite statement, and it is shared at onboarding."
    _, answer, found = qs.answer_question(s, ws.id, question.id, text, None, MODELS, TODAY)
    assert answer is not None and found == []
    assert "Dana Ortiz" not in answer.text and "risk appetite statement" in answer.text
    o = csf.framework().get("GV.RM-02")
    assert csf.gap_label(o, answer.label, answer.value, answer.statement_id) == "confirmed_by_you"


def test_a_govern_answer_suggests_fills_for_govern_parts_only_until_accepted(s: Session) -> None:
    # GV.RM-02 (Risk Management Strategy) and GV.PO-01/02 (Policy) share only the CSF function Govern
    # (decision 8)
    ws, q, run, llm, found = _govern_fills(s)
    po1, po2 = _item(s, q.id, "GV.PO-01"), _item(s, q.id, "GV.PO-02")
    # one re-check per open Govern part: GV.PO-01's 3 and GV.PO-02's 4, under MAX_RECHECKS; nothing else
    assert [r.item_id for r in llm.requests] == [f"{po1.id}#{n}" for n in (1, 2, 3)] + [
        f"{po2.id}#{n}" for n in (1, 2, 3, 4)
    ]
    assert len(llm.requests) <= qs.MAX_RECHECKS
    assert (
        s.get_one(Answer, _answer_id(s, run.id, po1.id)).label == "unknown"
    )  # nothing applied until accepted
    assert sorted((sg.item_id == po1.id, sg.part) for sg in found) == [
        (False, 1),
        (False, 2),
        (False, 3),
        (False, 4),
        (True, 1),
        (True, 2),
        (True, 3),
    ]
    fill = next(sg for sg in found if sg.item_id == po1.id and sg.part == 2)
    a = qs.accept_suggestion(s, ws.id, fill.id)
    # Ruling 6: one filled part plus Gaps stays Partly covered, and the filled part is named as the visitor's
    assert (a.label, a.value, a.approved_at, a.statement_id) == ("partial", "Partial", None, None)
    assert a.text.startswith("Confirmed by you: part 2. No evidence: parts 1, 3.")
    assert a.citations[0]["filename"] == "answer-002.txt"  # the visitor's statement (GV.RM-02 is item 2)
    stored = s.scalars(select(RunItem).where(RunItem.run_id == run.id, RunItem.item_id == po1.id)).one().parts
    assert (stored["2"]["label"], stored["2"]["statement_id"]) == ("verified", str(fill.statement_id))
    assert "model" not in stored["2"] and "stance_prompt" not in stored["2"]  # no judge of its own
    states = dict(
        s.execute(select(SuggestedFill.part, SuggestedFill.status).where(SuggestedFill.item_id == po1.id))
        .tuples()
        .all()
    )
    assert states == {1: "open", 2: "accepted", 3: "open"}  # only that part's other fills are dismissed
    with pytest.raises(qs.Conflict):
        qs.accept_suggestion(s, ws.id, fill.id)
    for n in (1, 3):  # accepting one fill never locks the others (Ruling 6)
        a = qs.accept_suggestion(
            s, ws.id, next(sg.id for sg in found if sg.item_id == po1.id and sg.part == n)
        )
    # every part Covered, from the visitor's answer: Confirmed by you, never Covered (adversary-1 I3)
    assert (a.label, a.statement_id) == ("user_confirmed", fill.statement_id)
    o = csf.framework().get("GV.PO-01")
    assert csf.gap_label(o, a.label, a.value, a.statement_id) == "confirmed_by_you"
    assert a.text.startswith("Confirmed by you: parts 1, 2, 3.")


def test_the_gap_sheet_marks_a_quote_from_the_visitors_answer(s: Session) -> None:
    ws, q, run, llm, found = _govern_fills(s)
    po1 = _item(s, q.id, "GV.PO-01")
    qs.accept_suggestion(s, ws.id, next(sg.id for sg in found if sg.item_id == po1.id and sg.part == 2))
    book = openpyxl.Workbook()
    export._write_gap(book.active, gap_sheet(s, s.get_one(Questionnaire, q.id), run))
    buf = io.BytesIO()
    book.save(buf)
    sheet = openpyxl.load_workbook(buf).active
    row = next(r for r in range(3, sheet.max_row + 1) if sheet.cell(r, 1).value == "GV.PO-01")
    assert sheet.cell(row, 6).value == f'"{SAID}" (answer-002.txt line 1) (your answer)'


def test_a_part_fill_is_shown_with_the_parts_wording(s: Session) -> None:
    ws, q, run, llm, found = _govern_fills(s)
    po2 = _item(s, q.id, "GV.PO-02")
    out = _suggestion_out(s, next(sg for sg in found if sg.item_id == po2.id and sg.part == 3))
    assert (out.part, out.question) == (3, csf.framework().get("GV.PO-02").parts[2])  # preflight I3


def test_accepting_a_fill_while_its_outcome_is_checked_again_is_a_conflict(s: Session) -> None:
    ws, q, run, llm, found = _govern_fills(s)
    po1 = _item(s, q.id, "GV.PO-01")
    s.execute(
        delete(Answer).where(Answer.run_id == run.id, Answer.item_id == po1.id)
    )  # as a re-open leaves it
    s.commit()
    with pytest.raises(qs.Conflict):  # adversary-1 M1: not a 404 "gone"
        qs.accept_suggestion(s, ws.id, next(sg.id for sg in found if sg.item_id == po1.id))


def test_an_answer_never_fills_another_functions_parts(s: Session) -> None:
    ws, q, run = _core_done(s)
    govern = [_item(s, q.id, c).id for c in ("GV.PO-01", "GV.PO-02")]
    s.execute(
        update(Answer)
        .where(Answer.run_id == run.id, Answer.item_id.in_(govern))
        .values(approved_at=func.now())
    )
    s.commit()  # Govern has no open Checked part left; every other function's outcomes are open Gaps
    question = next(x for x in qs.ensure_questions(s, ws.id, run.id) if _code(s, x.item_ids) == "GV.RM-02")
    llm = ByStepLLM({"recheck": REPLY})
    _, answer, found = qs.answer_question(s, ws.id, question.id, SAID, llm, MODELS, TODAY)
    assert answer is not None and answer.label == "user_confirmed"
    assert (llm.requests, found) == ([], [])  # no Protect, Detect, Identify, Respond or Recover part is asked

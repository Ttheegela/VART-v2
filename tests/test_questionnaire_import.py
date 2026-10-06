import io
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.api.schemas import Mapping
from app.ingest.parse import IngestError
from app.questionnaires import MAX_ITEMS, SAMPLE_DIR, detect, item_inputs, read_items, read_sheets
from evals import pack as packs

ROOT = Path(__file__).resolve().parent.parent
MAPPER = ROOT / "data" / "mapper"
EXPECTED: dict[str, dict[str, object]] = json.loads((MAPPER / "expected.json").read_text())


def _expected(raw: dict[str, object]) -> dict[str, object]:
    """expected.json writes csv columns as 1-based numbers and header_row as "1"; the API uses letters and
    ints for both formats (carry-over: mapper scoring notes)."""
    out = dict(raw)
    out["header_row"] = int(str(raw["header_row"]))
    for key in ("id_col", "question_col", "answer_col", "comments_col"):
        v = raw[key]
        out[key] = get_column_letter(int(v)) if isinstance(v, str) and v.isdigit() else v
    return out


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_the_mapper_finds_every_variant(name: str) -> None:
    # Spec 8: the column-mapping gate, 10/10 on the mapper set (moved to Plan 3 with the mapper).
    fmt, sheets = read_sheets(name, (MAPPER / name).read_bytes())
    found = detect(sheets, fmt)
    assert found is not None
    got = found.model_dump(exclude={"topic_col", "scope"})
    assert got == _expected(EXPECTED[name])


def test_every_variant_yields_its_twenty_questions_on_one_line() -> None:
    for name in EXPECTED:
        fmt, sheets = read_sheets(name, (MAPPER / name).read_bytes())
        items = read_items(sheets, detect(sheets, fmt))  # type: ignore[arg-type]
        assert len(items) == 20, name
        assert all("\n" not in i.question and i.question == " ".join(i.question.split()) for i in items)


@pytest.mark.parametrize(("name", "questionnaire"), [("vsq-a.xlsx", "vsq-a"), ("mvsp-b.csv", "mvsp-b")])
def test_the_samples_read_exactly_like_the_eval_items(name: str, questionnaire: str) -> None:
    # The dev recordings key on question and topic: the sample run (and Plan 4's precomputed path) replays
    # them only if the importer reads the bundled files exactly as the eval's selection does.
    fmt, sheets = read_sheets(name, (SAMPLE_DIR / name).read_bytes())
    mapping = detect(sheets, fmt)
    assert mapping is not None
    assert item_inputs(read_items(sheets, mapping)) == packs.load("dev").items(questionnaire)


def test_vsq_a_is_detected_as_its_committed_mapping() -> None:
    fmt, sheets = read_sheets("vsq-a.xlsx", (SAMPLE_DIR / "vsq-a.xlsx").read_bytes())
    committed = json.loads((SAMPLE_DIR / "vsq-a.mapping.json").read_text())
    assert detect(sheets, fmt) == Mapping(**committed, topic_col="B")


def test_more_than_150_items_is_refused() -> None:
    rows = "Question\n" + "".join(f"Is control {n} in place?\n" for n in range(MAX_ITEMS + 1))
    fmt, sheets = read_sheets("big.csv", rows.encode())
    with pytest.raises(IngestError, match="150"):
        read_items(sheets, detect(sheets, fmt))  # type: ignore[arg-type]


def test_a_zip_bomb_is_refused() -> None:
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", b"<x/>")
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (60 * 1024 * 1024))
    with pytest.raises(IngestError, match="too large once unpacked"):
        read_sheets("bomb.xlsx", buf.getvalue())


def test_only_xlsx_and_csv_are_questionnaires() -> None:
    with pytest.raises(IngestError, match="xlsx or csv"):
        read_sheets("policy.md", b"# Policy\n")


def test_a_file_with_no_question_column_has_no_mapping() -> None:
    wb = openpyxl.Workbook()
    wb.active.append(["Weight", "Owner"])
    wb.active.append([3, "IT"])
    buf = io.BytesIO()
    wb.save(buf)
    fmt, sheets = read_sheets("plain.xlsx", buf.getvalue())
    assert detect(sheets, fmt) is None

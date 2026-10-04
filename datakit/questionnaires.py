"""Build the bundled questionnaires from their selection files. `python -m datakit.questionnaires`"""

import csv
import html
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from datakit.schemas import Selection, load_yaml

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "data" / "sources"
OUT = ROOT / "data" / "questionnaires"
FIXED = datetime(2026, 1, 1)
QUESTION_TYPES = {"yesno", "radiogroup", "checkgroup", "line", "box"}
_TAG = re.compile(r"<[^>]+>")
_CONCAT = re.compile(r'"\s*\+\s*\n\s*"')  # VSAQ files join long strings JavaScript-style: "..." +\n "..."
_STAMP = re.compile(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")
_GENERATOR = re.compile(rb"<(Application|AppVersion)>[^<]*</\1>")  # app.xml names the tool and its version
MAPPING: dict[str, str | int] = {
    "sheet": "Questionnaire",
    "header_row": 5,
    "id_col": "A",
    "question_col": "C",
    "answer_col": "D",
    "comments_col": "E",
}


def _plain(text: str) -> str:
    return " ".join(html.unescape(_TAG.sub(" ", text)).split())


def _walk(node: Any, found: dict[str, str], file: str) -> None:
    if isinstance(node, dict):
        if node.get("type") in QUESTION_TYPES and node.get("id") and node.get("text"):
            found[f"vsaq:{file}#{node['id']}"] = _plain(node["text"])
        for value in node.values():
            _walk(value, found, file)
    elif isinstance(node, list):
        for value in node:
            _walk(value, found, file)


def vsaq_items() -> dict[str, str]:
    found: dict[str, str] = {}
    for path in sorted((SOURCES / "vsaq").glob("*.json")):
        body = "\n".join(
            line
            for line in path.read_text(encoding="utf-8").splitlines()
            if not line.lstrip().startswith("//")
        )
        _walk(json.loads(_CONCAT.sub("", body)), found, path.stem)
    return found


def mvsp_items() -> dict[str, str]:
    catalog = json.loads((SOURCES / "mvsp" / "mvsp_v3.0-20231109-catalog.json").read_text(encoding="utf-8"))
    found: dict[str, str] = {}
    for group in catalog["catalog"]["groups"]:
        for control in group.get("controls", []):
            label = next(p["value"] for p in control.get("props", []) if p["name"] == "label")
            found[f"mvsp:{label}"] = control["title"]
    return found


def normalize_zip(path: Path) -> None:
    """Make an Office zip byte-identical on every rebuild, on any machine: fixed timestamps, stored (not
    deflated) members so no zlib build can differ, one creating OS (CPython records 0 on Windows and 3
    elsewhere) and no tool name or version in docProps/app.xml."""
    with zipfile.ZipFile(path) as zin:
        parts = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zout:
        for name, data in parts:
            if name == "docProps/core.xml":
                data = _STAMP.sub(rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
            elif name == "docProps/app.xml":
                data = _GENERATOR.sub(b"", data)
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED
            info.create_system = 3
            info.external_attr = 0o644 << 16
            zout.writestr(info, data)


def build_xlsx(sel: Selection, path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED
    intro = wb.active
    intro.title = "Instructions"
    intro["A1"] = f"{sel.buyer} - {sel.title}"
    intro["A1"].font = Font(bold=True, size=14)
    for row, text in enumerate(
        [
            "Please answer every question in the Questionnaire sheet.",
            "Use Yes, No or N/A in the Response column and give evidence or context in Comments.",
            "Return the completed file to your Northbeam Health contact.",
        ],
        start=3,
    ):
        intro.cell(row, 1, text)
    ws = wb.create_sheet("Questionnaire")
    ws["A1"] = sel.title
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Vendor name:"
    ws["A3"] = "Completed by:"
    headers = ["#", "Domain", "Control Question", "Response (Yes / No / N/A)", "Comments / Evidence"]
    for col, text in enumerate(headers, start=1):
        ws.cell(5, col, text).font = Font(bold=True)
    validation = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    ws.add_data_validation(validation)
    row, section = 6, None
    for item in sel.items:
        if item.section != section:
            section = item.section
            ws.cell(row, 1, section.upper()).font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            row += 1
        ws.cell(row, 1, item.code)
        ws.cell(row, 2, item.section)
        ws.cell(row, 3, item.question).alignment = Alignment(wrap_text=True, vertical="top")
        validation.add(f"D{row}")
        row += 1
    for letter, width in zip("ABCDE", (10, 22, 70, 18, 40), strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A6"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    normalize_zip(path)


def build_csv(sel: Selection, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["ID", "Control", "Question", "Answer", "Notes"])
        for item in sel.items:
            w.writerow([item.code, item.section, item.question, "", ""])


def mapping_json() -> str:
    return json.dumps(MAPPING, indent=2) + "\n"


def main() -> None:
    build_xlsx(load_yaml(OUT / "vsq-a.selection.yaml", Selection), OUT / "vsq-a.xlsx")
    build_csv(load_yaml(OUT / "mvsp-b.selection.yaml", Selection), OUT / "mvsp-b.csv")
    (OUT / "vsq-a.mapping.json").write_text(mapping_json(), encoding="utf-8", newline="\n")


if __name__ == "__main__":
    main()

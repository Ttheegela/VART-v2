"""Ten deliberately messy variants of the first 20 VSQ-A items, each with the column mapping a mapper should
find (Plan 2's column mapper is tested against them).   python -m datakit.mapper_variants

data/mapper/expected.json maps every file name to {sheet, header_row, id_col, question_col, answer_col,
comments_col}. In a workbook the sheet is its name, header_row a row number and the columns are letters. In a
CSV file the sheet is null, header_row is "1" and the columns are 1-based numbers as strings. id_col is null
where the file has no ID column. Where the header takes two rows (v04), header_row is the row that holds
"Question".

Byte-identical on every run and every machine: workbooks go through normalize_zip (fixed timestamps, stored
members, no tool name); CSV files are UTF-8 without a BOM and end every record with CRLF (RFC 4180, as Excel
writes it), set below rather than left to the csv module's default.
"""

import csv
import json
import textwrap
from collections import Counter
from collections.abc import Callable
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font

from datakit.questionnaires import FIXED, normalize_zip
from datakit.questionnaires import OUT as QDIR
from datakit.schemas import Selection, SelectionItem, load_yaml

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "mapper"
Variant = Callable[[list[SelectionItem], Path], dict[str, Any]]


def _workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED
    return wb


def _save(wb: openpyxl.Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    normalize_zip(path)


def _csv(path: Path, rows: list[list[str]], delimiter: str = ",") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f, delimiter=delimiter, lineterminator="\r\n").writerows(rows)


def _mapping(
    sheet: str | None,
    header_row: int | str,
    *,
    id_col: str | None,
    question_col: str,
    answer_col: str,
    comments_col: str,
) -> dict[str, Any]:
    return {
        "sheet": sheet,
        "header_row": header_row,
        "id_col": id_col,
        "question_col": question_col,
        "answer_col": answer_col,
        "comments_col": comments_col,
    }


def v01(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Clean: the header is row 1 of the only sheet, and there is no ID column."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Question", "Answer", "Comments"])
    for item in items:
        ws.append([item.question])
    _save(wb, path)
    return _mapping("Sheet1", 1, id_col=None, question_col="A", answer_col="B", comments_col="C")


def v02(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """A title block above the table: the header is on row 6 and the ID column comes first."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Questionnaire"
    ws["A1"] = "Northbeam Health - Supplier Security Review"
    ws["A2"] = "Please complete every row."
    ws["A4"] = "Version 3"
    for col, text in enumerate(["ID", "Question", "Response", "Notes"], start=1):
        ws.cell(6, col, text)
    for item in items:
        ws.append([item.code, item.question])
    _save(wb, path)
    return _mapping("Questionnaire", 6, id_col="A", question_col="B", answer_col="C", comments_col="D")


def v03(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Two sheets: a text-only Cover first, then Security, whose header is row 2 under a banner."""
    wb = _workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["A1"] = "Northbeam Health - Supplier Security Review"
    cover["A3"] = "Please answer every question on the Security sheet."
    cover["A4"] = "Return the completed workbook to your Northbeam Health contact."
    ws = wb.create_sheet("Security")
    ws["A1"] = "Security controls"
    ws.append(["Ref", "Questions", "Answer", "Comments"])
    for item in items:
        ws.append([item.code, item.question])
    _save(wb, path)
    return _mapping("Security", 2, id_col="A", question_col="B", answer_col="C", comments_col="D")


def v04(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """A two-row header: Question (merged down), Response (merged over Yes/No and Details); data at row 5."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Assessment"
    ws["A1"] = "Supplier Security Assessment"
    ws["A3"] = "Question"
    ws["B3"] = "Response"
    ws["B4"] = "Yes/No"
    ws["C4"] = "Details"
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:C3")
    for item in items:
        ws.append([item.question])
    _save(wb, path)
    return _mapping("Assessment", 3, id_col=None, question_col="A", answer_col="B", comments_col="C")


def v05(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Merged, bold section rows between the questions, and a blank row between sections."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Controls"
    ws.append(["#", "Control Question", "Response", "Notes"])
    row, section = 2, None
    for item in items:
        if item.section != section:
            if section is not None:
                row += 1
            section = item.section
            ws.cell(row, 1, section.upper()).font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            row += 1
        ws.cell(row, 1, item.code)
        ws.cell(row, 2, item.question)
        row += 1
    _save(wb, path)
    return _mapping("Controls", 1, id_col="A", question_col="B", answer_col="C", comments_col="D")


# The vendor's replies to the first five items in v06: (answer, comment).
FILLED = [
    ("Yes", "See the information security policy."),
    ("Yes", None),
    ("Yes", "Named in the policy."),
    ("No", "Not certified."),
    ("Yes", "Report available on request."),
]


def v06(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """The answers (and some comments) are already filled in for the first five rows."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(["Question", "Answer", "Comments"])
    for n, item in enumerate(items):
        ws.append([item.question, *(FILLED[n] if n < len(FILLED) else ())])
    _save(wb, path)
    return _mapping("Responses", 1, id_col=None, question_col="A", answer_col="B", comments_col="C")


def v07(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Semicolon-delimited, as a spreadsheet in a comma-decimal locale exports it."""
    _csv(path, [["Question", "Answer", "Comment"], *([i.question, "", ""] for i in items)], delimiter=";")
    return _mapping(None, "1", id_col=None, question_col="1", answer_col="2", comments_col="3")


def v08(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Comma-delimited with the longer questions broken over several lines inside quotes; the ID column is
    headed "Question ID", which is not the question column."""
    rows = [["Question ID", "Question", "Answer", "Comments"]]
    rows += [[i.code, textwrap.fill(i.question, 70), "", ""] for i in items]
    _csv(path, rows)
    return _mapping(None, "1", id_col="1", question_col="2", answer_col="3", comments_col="4")


def v09(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """Spanish headers over the English questions."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Cuestionario"
    ws.append(["Pregunta", "Respuesta", "Comentarios"])
    for item in items:
        ws.append([item.question])
    _save(wb, path)
    return _mapping("Cuestionario", 1, id_col=None, question_col="A", answer_col="B", comments_col="C")


PREFIX = {"Governance": "GV", "Access Control": "AC", "Data Security": "DS"}
WEIGHTS = (3, 5, 2, 1, 4)
OWNERS = ("Security", "IT Ops", "Compliance", "Procurement")


def v10(items: list[SelectionItem], path: Path) -> dict[str, Any]:
    """The question is the third column, between the IDs (AC-01, not VSQ-01) and the answer; extra Weight and
    Owner columns sit before the IDs and after the comments."""
    wb = _workbook()
    ws = wb.active
    ws.title = "Scorecard"
    ws.append(["Weight", "Control ID", "Question", "Answer", "Comments", "Owner"])
    count: Counter[str] = Counter()
    for n, item in enumerate(items):
        prefix = PREFIX[item.section]
        count[prefix] += 1
        weight, owner = WEIGHTS[n % len(WEIGHTS)], OWNERS[n % len(OWNERS)]
        ws.append([weight, f"{prefix}-{count[prefix]:02d}", item.question, None, None, owner])
    _save(wb, path)
    return _mapping("Scorecard", 1, id_col="B", question_col="C", answer_col="D", comments_col="E")


VARIANTS: list[Variant] = [v01, v02, v03, v04, v05, v06, v07, v08, v09, v10]
CSV = ("v07", "v08")  # the other eight are workbooks


def build_all(out: Path) -> dict[str, dict[str, Any]]:
    """Write the ten files and expected.json into `out`; return each file name's expected mapping."""
    items = list(load_yaml(QDIR / "vsq-a.selection.yaml", Selection).items[:20])
    expected: dict[str, dict[str, Any]] = {}
    for build in VARIANTS:
        name = f"{build.__name__}.{'csv' if build.__name__ in CSV else 'xlsx'}"
        expected[name] = build(items, out / name)
    text = json.dumps(expected, indent=2, sort_keys=True) + "\n"
    (out / "expected.json").write_text(text, encoding="utf-8", newline="\n")
    return expected


def main() -> None:
    build_all(OUT)


if __name__ == "__main__":
    main()

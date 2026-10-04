"""Read a rendered document back as text lines, the way a visitor's upload would be read. This is the
reference used to validate the dev pack; Plan 2's ingest parser is tested against it.

Line rules (spec section 6.4): one paragraph, heading, list item or table row per line; table rows and
spreadsheet rows use app.text.record_line; PDF text is joined across visual lines because wrapping is not
a paragraph break."""

import re
from pathlib import Path

import docx
import openpyxl
import pypdfium2
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.text import normalize, record_line

_TABLE_RULE = re.compile(r"^\|?\s*:?-{3,}")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _md_lines(text: str) -> list[str]:
    out: list[str] = []
    para: list[str] = []
    header: list[str] | None = None

    def flush() -> None:
        if para:
            out.append(normalize(" ".join(para)))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            header = None
            continue
        if line.lstrip().startswith("|"):
            flush()
            if header is None:
                header = _cells(line)
            elif not _TABLE_RULE.match(line.strip()):
                out.append(record_line(header, _cells(line)))
            continue
        if line.startswith("#"):
            flush()
            out.append(normalize(line.lstrip("#")))
        elif line.lstrip().startswith(("- ", "* ")):
            flush()
            out.append(normalize(line.lstrip()[2:]))
        else:
            para.append(line)
    flush()
    return [x for x in out if x]


def _docx_lines(path: Path) -> list[str]:
    out: list[str] = []
    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                out.append(normalize(block.text))
        elif isinstance(block, Table):
            rows = [[c.text for c in r.cells] for r in block.rows]
            if rows:
                out.extend(record_line(rows[0], r) for r in rows[1:])
    return [x for x in out if x]


def _xlsx_lines(path: Path) -> list[str]:
    out: list[str] = []
    wb = openpyxl.load_workbook(str(path), data_only=True)
    for ws in wb.worksheets:
        header: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            values = list(row)
            filled = [v for v in values if v not in (None, "")]
            if not filled:
                continue
            if header is None and len(filled) >= 2 and all(isinstance(v, str) for v in filled):
                header = [str(v) if v is not None else "" for v in values]
                continue
            if header is None:
                out.append(normalize(" ".join(str(v) for v in filled)))
            else:
                out.append(record_line(header, values))
    return [x for x in out if x]


def _pdf_text(path: Path) -> str:
    pdf = pypdfium2.PdfDocument(str(path))
    try:
        return normalize(" ".join(page.get_textpage().get_text_range() for page in pdf))
    finally:
        pdf.close()


def lines_of(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".md":
        return _md_lines(path.read_text(encoding="utf-8"))
    if suffix == ".docx":
        return _docx_lines(path)
    if suffix == ".xlsx":
        return _xlsx_lines(path)
    if suffix == ".pdf":
        return [_pdf_text(path)]
    raise ValueError(f"unsupported file type: {path.name}")


def text_of(path: Path) -> str:
    return "\n".join(lines_of(path))
